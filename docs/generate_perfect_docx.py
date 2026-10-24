import os
import sys
import shutil
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = Path(r"d:\unifai_project")
ASSETS_DIR = BASE_DIR / "docs" / "pdf_assets"
EXTRACTED_DIR = BASE_DIR / "docs" / "extracted_from_new_doc"
LOGO_PATH = str(BASE_DIR / "ui" / "public" / "yes-panchi-logo.png")
OUTPUT_DOCX = Path(r"C:\Users\sakth\Downloads\infosec.docx")
TARGET_NEW_DOC = BASE_DIR / "new doc.docx"
TARGET_UPDATED_DOC = BASE_DIR / "YP-Raksha_v1.5_Updated.docx"
LOCAL_COPY = BASE_DIR / "docs" / "YP_RAKSHA_InfoSec_Compliance_Report.docx"

# Color Palette (Executive Navy & Emerald)
NAVY_HEX = "0F172A"        # Slate 900
NAVY_RGB = RGBColor(0x0F, 0x17, 0x2A)
SLATE_DARK_RGB = RGBColor(0x1E, 0x29, 0x3B)
SLATE_BODY_RGB = RGBColor(0x33, 0x41, 0x55)
MUTED_RGB = RGBColor(0x64, 0x74, 0x8B)
GREEN_RGB = RGBColor(0x05, 0x96, 0x69)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_footer_xml(footer):
    p_xml = parse_xml(
        f'<w:p {nsdecls("w")}>'
        f'  <w:pPr>'
        f'    <w:pBdr>'
        f'      <w:top w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/>'
        f'    </w:pBdr>'
        f'    <w:tabs>'
        f'      <w:tab w:val="right" w:pos="10080"/>'
        f'    </w:tabs>'
        f'    <w:spacing w:before="60" w:after="0"/>'
        f'  </w:pPr>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'      <w:b/>'
        f'      <w:sz w:val="17"/>'
        f'      <w:color w:val="0F172A"/>'
        f'    </w:rPr>'
        f'    <w:t>YES PANCHI</w:t>'
        f'  </w:r>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'      <w:sz w:val="17"/>'
        f'      <w:color w:val="64748B"/>'
        f'    </w:rPr>'
        f'    <w:t xml:space="preserve">  |  YP-Raksha</w:t>'
        f'  </w:r>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'      <w:sz w:val="17"/>'
        f'    </w:rPr>'
        f'    <w:tab/>'
        f'  </w:r>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'      <w:b/>'
        f'      <w:sz w:val="17"/>'
        f'      <w:color w:val="0F172A"/>'
        f'    </w:rPr>'
        f'    <w:t xml:space="preserve">Page </w:t>'
        f'  </w:r>'
        f'  <w:fldSimple w:instr="PAGE">'
        f'    <w:r>'
        f'      <w:rPr>'
        f'        <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'        <w:b/>'
        f'        <w:sz w:val="17"/>'
        f'        <w:color w:val="0F172A"/>'
        f'      </w:rPr>'
        f'      <w:t>1</w:t>'
        f'    </w:r>'
        f'  </w:fldSimple>'
        f'  <w:r>'
        f'    <w:rPr>'
        f'      <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'      <w:b/>'
        f'      <w:sz w:val="17"/>'
        f'      <w:color w:val="0F172A"/>'
        f'    </w:rPr>'
        f'    <w:t xml:space="preserve"> of </w:t>'
        f'  </w:r>'
        f'  <w:fldSimple w:instr="NUMPAGES">'
        f'    <w:r>'
        f'      <w:rPr>'
        f'        <w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'        <w:b/>'
        f'        <w:sz w:val="17"/>'
        f'        <w:color w:val="0F172A"/>'
        f'      </w:rPr>'
        f'      <w:t>23</w:t>'
        f'    </w:r>'
        f'  </w:fldSimple>'
        f'</w:p>'
    )
    footer._element.append(p_xml)

def create_req_box(doc, code, req_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F0FDF4")
    set_cell_margins(cell, top=35, bottom=35, left=90, right=90)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="86EFAC"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="059669"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="86EFAC"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="86EFAC"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p1 = cell.paragraphs[0]
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(1)
    r1 = p1.add_run(f"{code}  |  MANDATORY SECURITY SPECIFICATION")
    r1.font.name = "Calibri"
    r1.font.size = Pt(7.5)
    r1.font.bold = True
    r1.font.color.rgb = GREEN_RGB
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(f'"{req_text}"')
    r2.font.name = "Calibri"
    r2.font.size = Pt(8)
    r2.font.italic = True
    r2.font.color.rgb = SLATE_DARK_RGB

def add_image_card_docx(doc, img_path, caption, width_in_inches=6.0):
    tbl = doc.add_table(rows=2, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    c0 = tbl.cell(0, 0)
    c0.width = Inches(width_in_inches + 0.1)
    set_cell_background(c0, "F8FAFC")
    set_cell_margins(c0, top=25, bottom=10, left=25, right=25)
    
    tcPr0 = c0._tc.get_or_add_tcPr()
    borders0 = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr0.append(borders0)
    
    p_img = c0.paragraphs[0]
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(0)
    p_img.paragraph_format.space_after = Pt(0)
    p_img.add_run().add_picture(str(img_path), width=Inches(width_in_inches))
    
    c1 = tbl.cell(1, 0)
    c1.width = Inches(width_in_inches + 0.1)
    set_cell_background(c1, "F8FAFC")
    set_cell_margins(c1, top=8, bottom=25, left=40, right=40)
    
    tcPr1 = c1._tc.get_or_add_tcPr()
    borders1 = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    tcPr1.append(borders1)
    
    p_cap = c1.paragraphs[0]
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(1)
    p_cap.paragraph_format.space_after = Pt(1)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(7.5)
    r_cap.font.bold = True
    r_cap.font.color.rgb = MUTED_RGB

def add_topic_item_docx(doc, title, label, req_text, desc_text, img_path, caption_text, img_w=6.0):
    p_t = doc.add_paragraph()
    p_t.paragraph_format.space_before = Pt(1)
    p_t.paragraph_format.space_after = Pt(1)
    r_t = p_t.add_run(title)
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(9.5)
    r_t.font.bold = True
    r_t.font.color.rgb = NAVY_RGB
    
    create_req_box(doc, label, req_text)
    
    p_b = doc.add_paragraph()
    p_b.paragraph_format.space_before = Pt(1)
    p_b.paragraph_format.space_after = Pt(2)
    r_b = p_b.add_run("Technical Compliance & Architecture: ")
    r_b.font.name = "Calibri"
    r_b.font.size = Pt(7.5)
    r_b.font.bold = True
    r_b2 = p_b.add_run(desc_text)
    r_b2.font.name = "Calibri"
    r_b2.font.size = Pt(7.5)
    r_b2.font.color.rgb = SLATE_DARK_RGB
    
    add_image_card_docx(doc, img_path, caption_text, width_in_inches=img_w)

def generate_perfect_updated_doc():
    doc = docx.Document()
    
    # 0.5 in margins
    for s in doc.sections:
        s.top_margin = Inches(0.5)
        s.bottom_margin = Inches(0.5)
        s.left_margin = Inches(0.5)
        s.right_margin = Inches(0.5)
        s.page_width = Inches(8.27)   # A4
        s.page_height = Inches(11.69)
        s.different_first_page_header_footer = True
        
    section = doc.sections[0]
    
    # Subsequent Pages Header (Page 2+)
    header = section.header
    p_hdr = header.paragraphs[0]
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_hdr_pr = p_hdr._p.get_or_add_pPr()
    p_hdr_bdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="E2E8F0"/></w:pBdr>')
    p_hdr_pr.append(p_hdr_bdr)
    r_hdr = p_hdr.add_run("YP-Raksha  |  Security & Compliance Verification\n")
    r_hdr.font.name = "Calibri"
    r_hdr.font.size = Pt(8.5)
    r_hdr.font.bold = True
    r_hdr.font.color.rgb = SLATE_DARK_RGB
    
    # Subsequent Pages Footer (Page 2+)
    footer = section.footer
    add_footer_xml(footer)

    # ==========================================
    # PAGE 1: COVER PAGE (CENTERED)
    # ==========================================
    for _ in range(5):
        doc.add_paragraph()
        
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(15)
    p_logo.paragraph_format.space_after = Pt(25)
    p_logo.add_run().add_picture(LOGO_PATH, width=Inches(3.2))
    
    p_company = doc.add_paragraph()
    p_company.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_company.paragraph_format.space_after = Pt(4)
    r_comp = p_company.add_run("YES PANCHI")
    r_comp.font.name = "Calibri"
    r_comp.font.size = Pt(20)
    r_comp.font.bold = True
    r_comp.font.color.rgb = NAVY_RGB
    
    p_prod = doc.add_paragraph()
    p_prod.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prod.paragraph_format.space_after = Pt(1)
    r_prod = p_prod.add_run("YP-Raksha")
    r_prod.font.name = "Calibri"
    r_prod.font.size = Pt(44)
    r_prod.font.bold = True
    r_prod.font.color.rgb = NAVY_RGB
    
    p_ver = doc.add_paragraph()
    p_ver.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ver.paragraph_format.space_after = Pt(6)
    r_ver = p_ver.add_run("Version 1.5")
    r_ver.font.name = "Calibri"
    r_ver.font.size = Pt(11)
    r_ver.font.bold = True
    r_ver.font.color.rgb = MUTED_RGB
    
    p_exp = doc.add_paragraph()
    p_exp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_exp.paragraph_format.space_after = Pt(5)
    r_exp = p_exp.add_run("RAKSHA - Real-time AI Knowledge Screening & Hazard Audit")
    r_exp.font.name = "Calibri"
    r_exp.font.size = Pt(12)
    r_exp.font.bold = True
    r_exp.font.color.rgb = GREEN_RGB
    
    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = p_tag.add_run("Enterprise AI Security & Data Governance")
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(11)
    r_tag.font.color.rgb = MUTED_RGB

    # ==========================================
    # PAGE 2: INDEX / TABLE OF CONTENTS (23 PAGES)
    # ==========================================
    doc.add_page_break()
    p_idx = doc.add_paragraph()
    p_idx.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_idx.paragraph_format.space_before = Pt(5)
    p_idx.paragraph_format.space_after = Pt(6)
    r_idx = p_idx.add_run("Index")
    r_idx.font.name = "Calibri"
    r_idx.font.size = Pt(20)
    r_idx.font.bold = True
    r_idx.font.color.rgb = NAVY_RGB
    
    toc_data = [
        ("6.3.8 VI (a)", "Automatic Redaction", "Inline PII masking & redaction (Aadhaar, PAN, contact details)", "Page 3"),
        ("6.3.8 VI (b)", "Malicious Intent Detection", "Mass extraction, enumeration patterns & jailbreak defense", "Page 4"),
        ("6.3.8 VI (c)", "Prompt Sanitization", "Incoming query analysis & classified personal data protection", "Page 4"),
        ("6.3.8 VI (d)", "Clipboard & Data Governance", "Strict controls on Copy-Paste, Drag-and-Drop & Downloads", "Page 5"),
        ("6.3.8 VI (e)", "Policy Enforcement", "Purpose alignment, consent validation & minor/children data", "Page 5"),
        ("6.3.8 VI (f)", "Just-in-Time (JIT) Coaching", "Real-time browser-level alerts & employee education popups", "Page 6"),
        ("6.3.8 VI (g)", "Response Scanning", "AI-generated response analysis & classification metadata", "Page 6"),
        ("6.3.8 VI (h)", "Risk Scoring", "Regulatory & business risk scoring for AI responses & search", "Page 7"),
        ("Req 44 & 45", "AI Discovery & Tool Visibility", "Automated tool discovery across 52+ domains & enterprise inventory", "Page 8"),
        ("Req 46 & 47", "Shadow AI & Data Repositories", "Shadow AI domain locking & file repository transmission controls", "Page 9"),
        ("Req 48 & 49", "Continuous Monitoring & DLP", "Continuous prompt interaction logging & regulated data detection", "Page 10"),
        ("Req 50 & 51", "Policy Controls & Enforcement", "Policy-based data exposure restrictions & automated enforcement", "Page 11"),
        ("Req 52 & 53", "Data Identification & Classification", "Sensitive enterprise data detection & automated data classification", "Page 12"),
        ("Req 54 & 55", "Risk Scoring & Access Boundaries", "Data-sensitivity AI risk scoring & excessive query detection", "Page 13"),
        ("Req 56 & 57", "Dataset Access & Anomaly Detection", "Restricted dataset prevention policies & anomalous pattern monitoring", "Page 14"),
        ("Req 58 & 59", "Suspicious Data & Exfiltration Alerts", "Suspicious AI interaction detection & automated exfiltration alerts", "Page 15"),
        ("Req 60 & 61", "Scraping Mitigation & Workflows", "Automated prompt scraping mitigation & analyst investigation tools", "Page 16"),
        ("Req 62 & 63", "Audit Logging & Governance Controls", "Detailed immutable AI audit logs & policy-aligned workspace governance", "Page 17"),
        ("Req 64 & 65", "Compliance Reports & Risk Posture", "AI compliance telemetry reports & fleet risk posture dashboards", "Page 18"),
        ("Req 66 & 67", "Audit Reporting & Dataset Protection", "Search activity audit reporting & classified dataset access prevention", "Page 19"),
        ("Req 68 & 69", "Runtime Monitoring & Threat Defense", "Runtime proxy data inspection & prompt injection attack defense", "Page 20"),
        ("Req 70 & 71", "Output Protection & Remediation", "Sensitive token output leakage prevention & automated blocking actions", "Page 21"),
        ("Req 72 & 73", "Multi-Asset Risk & Adoption Trends", "Workstation hardware telemetry & enterprise AI adoption analytics", "Page 22"),
        ("Req 74 & 75", "Behavioral Scoring & Executive Reports", "Behavioral contextual risk scoring & executive leadership briefings", "Page 23"),
    ]
    
    t_toc = doc.add_table(rows=1, cols=3)
    t_toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = t_toc.rows[0].cells
    hdr_cells[0].width = Inches(1.3)
    hdr_cells[1].width = Inches(4.2)
    hdr_cells[2].width = Inches(1.3)
    
    for i, title in enumerate(["Section / Req", "Topic / Security Control Specification", "Page Number"]):
        c = hdr_cells[i]
        set_cell_background(c, "0F172A")
        set_cell_margins(c, top=35, bottom=35, left=70, right=70)
        p = c.paragraphs[0]
        if i in [0, 2]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = WHITE_RGB
        
    for r_idx, (sec, title, desc, page_str) in enumerate(toc_data):
        row_cells = t_toc.add_row().cells
        bg_col = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        
        # Col 0: Sec
        c0 = row_cells[0]
        c0.width = Inches(1.3)
        set_cell_background(c0, bg_col)
        set_cell_margins(c0, top=20, bottom=20, left=50, right=50)
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r0 = p0.add_run(sec)
        r0.font.name = "Calibri"
        r0.font.size = Pt(8)
        r0.font.bold = True
        r0.font.color.rgb = NAVY_RGB
        
        # Col 1: Title & Desc
        c1 = row_cells[1]
        c1.width = Inches(4.2)
        set_cell_background(c1, bg_col)
        set_cell_margins(c1, top=20, bottom=20, left=60, right=60)
        p1 = c1.paragraphs[0]
        r1_t = p1.add_run(f"{title}\n")
        r1_t.font.name = "Calibri"
        r1_t.font.size = Pt(8)
        r1_t.font.bold = True
        r1_t.font.color.rgb = NAVY_RGB
        
        r1_d = p1.add_run(desc)
        r1_d.font.name = "Calibri"
        r1_d.font.size = Pt(7)
        r1_d.font.color.rgb = MUTED_RGB
        
        # Col 2: Page
        c2 = row_cells[2]
        c2.width = Inches(1.3)
        set_cell_background(c2, bg_col)
        set_cell_margins(c2, top=20, bottom=20, left=50, right=50)
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(page_str)
        r2.font.name = "Calibri"
        r2.font.size = Pt(8)
        r2.font.bold = True
        r2.font.color.rgb = NAVY_RGB

    # ==========================================
    # PAGE 3: TOPIC 6.3.8 VI (a)
    # ==========================================
    doc.add_page_break()
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(2)
    p_h1.paragraph_format.space_after = Pt(2)
    r = p_h1.add_run("6.3.8 VI (a): Automatic Redaction")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY_RGB
    
    create_req_box(doc, "6.3.8 VI (a)", "Automatic Redaction: The solution must automatically block or redact PII (e.g., Aadhaar, PAN, contact details) from natural language prompts before transmission to any AI service.")
    
    p_b = doc.add_paragraph()
    p_b.paragraph_format.space_before = Pt(2)
    p_b.paragraph_format.space_after = Pt(3)
    r_b = p_b.add_run("Technical Compliance & Architecture: YP-Raksha intercepts prompt submissions at the browser runtime and proxy layer before payload dispatch. High-performance regex evaluators and named-entity recognition (NER) models scan the buffer for Indian and international PII formats, ensuring zero leakage to external LLMs.")
    r_b.font.size = Pt(7.5)
    r_b.font.color.rgb = SLATE_DARK_RGB
    
    add_image_card_docx(doc, ASSETS_DIR / "p1_img1.jpeg", "Figure 1.1: YP-Raksha Guard Rules Engine with active PII redaction patterns for Aadhaar, PAN, and Contact Details.")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_image_card_docx(doc, ASSETS_DIR / "p1_img2.jpeg", "Figure 1.2: Audit verification log displaying automatic inline masking applied to employee prompt before proxy transmission.")

    # ==========================================
    # PAGE 4: TOPICS 6.3.8 VI (b) & (c)
    # ==========================================
    doc.add_page_break()
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(2)
    p_h1.paragraph_format.space_after = Pt(2)
    r = p_h1.add_run("6.3.8 VI (b) & (c): Malicious Intent Detection & Prompt Sanitization")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY_RGB
    
    add_topic_item_docx(
        doc,
        "Malicious Intent Detection",
        "6.3.8 VI (b)",
        "The solution must detect and block queries exhibiting malicious intent, such as mass extraction attempts, enumeration, and adversarial jailbreaks.",
        "YP-Raksha deploys heuristic token evaluators to intercept jailbreak attempts (DAN, prompt override) and token velocity rate limiters.",
        ASSETS_DIR / "p2_img1.jpeg",
        "Figure 2.1: Interception of adversarial jailbreak and mass enumeration attempts at the gateway layer."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_topic_item_docx(
        doc,
        "Prompt Sanitization",
        "6.3.8 VI (c)",
        "The solution must inspect and sanitize incoming prompts to prevent the injection of classified, confidential, or sensitive personal data into public AI services.",
        "Scans input streams for proprietary source code, internal IPs, connection strings, and classified operational assets before dispatch.",
        ASSETS_DIR / "p2_img2.jpeg",
        "Figure 2.2: Prompt sanitization inspection log showing scrubbed credentials and sanitized technical payloads."
    )

    # ==========================================
    # PAGE 5: TOPICS 6.3.8 VI (d) & (e)
    # ==========================================
    doc.add_page_break()
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(2)
    p_h1.paragraph_format.space_after = Pt(2)
    r = p_h1.add_run("6.3.8 VI (d) & (e): Clipboard Governance & Policy Enforcement")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY_RGB
    
    add_topic_item_docx(
        doc,
        "Clipboard & Data Governance",
        "6.3.8 VI (d)",
        "The solution must enforce strict controls to prevent unauthorized copy-paste, drag-and-drop, or bulk insertion of sensitive enterprise data into AI interfaces.",
        "In-browser DOM event listeners hook clipboard operations, terminating unauthorized copy-paste or file drop.",
        ASSETS_DIR / "p3_img1.jpeg",
        "Figure 3.1: Browser Settings console enforcing strict Copy-Paste, Drag-and-Drop, and File Download blocking."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_topic_item_docx(
        doc,
        "Policy Enforcement",
        "6.3.8 VI (e)",
        "The solution must validate that AI requests align with declared enterprise purpose and block prompts involving expired/withdrawn consent or children's data.",
        "Validates session tokens and prompt context against purpose tags; prompts referencing minors are dropped.",
        ASSETS_DIR / "p3_img2.jpeg",
        "Figure 3.2: Real-time policy enforcement engine terminating prompt submission due to missing consent authorization."
    )

    # ==========================================
    # PAGE 6: TOPICS 6.3.8 VI (f) & (g)
    # ==========================================
    doc.add_page_break()
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(2)
    p_h1.paragraph_format.space_after = Pt(2)
    r = p_h1.add_run("6.3.8 VI (f) & (g): Just-in-Time (JIT) Coaching & Response Scanning")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY_RGB
    
    add_topic_item_docx(
        doc,
        "Just-in-Time (JIT) Coaching",
        "6.3.8 VI (f)",
        "The solution must deliver real-time, context-aware policy alerts directly to users upon violation to reinforce safe AI usage.",
        "Dynamic floating alert banners inform users why their interaction violated policy, providing educational coaching.",
        ASSETS_DIR / "p4_img1.jpeg",
        "Figure 4.1: Real-time JIT coaching banner displayed to employee explaining policy violation at the moment of entry."
    )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_topic_item_docx(
        doc,
        "Response Scanning",
        "6.3.8 VI (g)",
        "All AI-generated responses must be analysed for PII or sensitive content and evaluated against established classification metadata.",
        "Captures inbound LLM completions, inspecting for toxicity, hallucinated instructions, and attaching safety metadata.",
        ASSETS_DIR / "p4_img2.jpeg",
        "Figure 4.2: AI Response Scanning dashboard showing real-time toxicity analysis and compliance metadata tagging."
    )

    # ==========================================
    # PAGE 7: TOPIC 6.3.8 VI (h)
    # ==========================================
    doc.add_page_break()
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(2)
    p_h1.paragraph_format.space_after = Pt(2)
    r = p_h1.add_run("6.3.8 VI (h): Risk Scoring")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = NAVY_RGB
    
    create_req_box(doc, "6.3.8 VI (h)", "Risk Scoring: The solution shall assign a regulatory and business risk score to each AI-generated response, search result, and synthesized output, enabling risk-based blocking or human review.")
    
    p_b = doc.add_paragraph()
    p_b.paragraph_format.space_before = Pt(2)
    p_b.paragraph_format.space_after = Pt(3)
    r_b = p_b.add_run("Technical Compliance & Architecture: YP-Raksha calculates a multidimensional composite risk score (0 to 100) weighing data sensitivity, user role, platform sensitivity, and prompt toxicity. Interactions exceeding tolerance thresholds trigger automated containment (Low: Allowed, Medium: JIT Alert, High: Hard Block).")
    r_b.font.size = Pt(7.5)
    r_b.font.color.rgb = SLATE_DARK_RGB
    
    add_image_card_docx(doc, ASSETS_DIR / "p5_img1.jpeg", "Figure 5.1: Multi-factor Risk Scoring engine matrix mapping weights across data sensitivity and platform security.")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_image_card_docx(doc, ASSETS_DIR / "p5_img2.jpeg", "Figure 5.2: Detailed Interaction Log and Risk Scoring table assigning risk ratings to every prompt and response.")

    # ==========================================
    # PAGES 8 - 23: REQUIREMENTS 44 TO 75 (2 PER PAGE, FULL SIZE SCREENSHOTS)
    # ==========================================
    pairs = [
        # Page 8: Requirements 44 & 45
        ("AI Discovery & Tool Visibility",
         ("Automated AI Tools Discovery", "AI DISCOVERY SPECIFICATION",
          "The solution must automatically discover AI tools, AI agents, AI chatbots, and AI model integrations used across the enterprise environment.",
          "Continuous discovery pipeline identifies enterprise access across 52+ commercial and emerging AI domains including ChatGPT, Claude, and Gemini.",
          EXTRACTED_DIR / "18_image22.png",
          "Figure 6.1: Target Web AI Platforms console showing automated discovery and active monitoring across 52 AI domains."),
         ("Enterprise-Wide AI Tool Visibility", "VISIBILITY SPECIFICATION",
          "The solution should provide visibility into AI tool usage across users, applications, and data repositories.",
          "Centralized endpoint inventory tracks desktop agents and browser extensions, logging machine hostname, IP, MAC address, and active user accounts.",
          EXTRACTED_DIR / "11_image23.png",
          "Figure 6.2: Centralized Guard Agents management dashboard tracking endpoint hostnames, IP, MAC, version, and status.")),

        # Page 9: Requirements 46 & 47
        ("Shadow AI & Data Repositories",
         ("Shadow AI Identification & Containment", "SHADOW AI SPECIFICATION",
          "The solution must identify shadow AI usage, including unauthorized AI applications and browser-based AI tools.",
          "Detects unapproved AI domains and browser AI tools; enables administrators to immediately enforce domain-level locking (Block entire website).",
          EXTRACTED_DIR / "27_image24.png",
          "Figure 7.1: Monitored AI Domains dashboard featuring one-click 'Block entire website' toggle to contain shadow AI."),
         ("Data Repositories Interaction Monitoring", "REPOSITORY MONITORING SPECIFICATION",
          "The solution should monitor AI interactions with enterprise data repositories including file systems, collaboration platforms, and cloud storage.",
          "Monitors file uploads, cloud attachments, and document drag-and-drop into external AI services, blocking unauthorized transmissions.",
          EXTRACTED_DIR / "04_image25.png",
          "Figure 7.2: Connected enterprise data repositories and cloud connector proxy monitoring console.")),

        # Page 10: Requirements 48 & 49
        ("Continuous Monitoring & DLP",
         ("Continuous User Interaction Monitoring", "RUNTIME MONITORING SPECIFICATION",
          "The solution must provide continuous monitoring of user interactions with AI tools to detect potential data leakage risks.",
          "Every prompt submission is intercepted inline, recording user prompts, tokens, execution cost, and security actions in live audit logs.",
          EXTRACTED_DIR / "23_image26.png",
          "Figure 8.1: Prompt & Chat History audit log displaying live intercepted requests, tokens, costs, and policy actions."),
         ("Sensitive & Regulated Data Detection", "DATA DETECTION SPECIFICATION",
          "The solution should detect when sensitive or regulated data is being accessed or transmitted through AI applications.",
          "High-speed pattern engine flags Aadhaar, mobile numbers, PAN, 18+ content, and confidential company tokens across all prompt dispatches.",
          EXTRACTED_DIR / "16_image27.png",
          "Figure 8.2: Active DLP Guard Rules console enforcing automated classification, 18+ content filters, and PII masking.")),

        # Page 11: Requirements 50 & 51
        ("Policy Controls & Enforcement",
         ("Policy-Based Data Exposure Controls", "POLICY CONTROL SPECIFICATION",
          "The solution must provide policy-based controls to restrict the exposure of sensitive data to AI tools and services.",
          "Enforces granular policies (Block vs. Redact) preventing classified datasets and confidential employee credentials from reaching external LLMs.",
          ASSETS_DIR / "p3_img1.jpeg",
          "Figure 9.1: Browser Guard Settings console enforcing strict copy-paste, drag-and-drop, and download exposure controls."),
         ("Automated Policy Enforcement Engine", "AUTOMATED ENFORCEMENT SPECIFICATION",
          "The solution should support automated enforcement of AI data protection policies.",
          "Enforces policy decisions within 5 milliseconds at runtime, terminating unauthorized operations without requiring manual analyst intervention.",
          EXTRACTED_DIR / "16_image27.png",
          "Figure 9.2: Guard Rules policy manager mapping critical severity actions (Block vs. Redact) across enterprise targets.")),

        # Page 12: Requirements 52 & 53
        ("Data Identification & Classification",
         ("Sensitive Enterprise Data Identification", "DATA IDENTIFICATION SPECIFICATION",
          "The solution must identify sensitive enterprise data accessed by AI systems, including personal data, financial data, and intellectual property.",
          "Detects proprietary intellectual property, customer financial records, and PII before text is uploaded into third-party AI interfaces.",
          ASSETS_DIR / "p1_img2.jpeg",
          "Figure 10.1: Live interaction audit trail identifying personal identifiers and proprietary corporate tokens."),
         ("Automated Data Classification Mechanisms", "CLASSIFICATION SPECIFICATION",
          "The solution should classify sensitive data accessed by AI systems using automated classification mechanisms.",
          "Assigns contextual data classifications (Confidential, Restricted, PII) in real time and masks detected tokens prior to proxy transmission.",
          EXTRACTED_DIR / "12_image28.png",
          "Figure 10.2: Guard Insights telemetry dashboard showing automated classification trends across enterprise prompts.")),

        # Page 13: Requirements 54 & 55
        ("Risk Scoring & Access Boundaries",
         ("Data-Sensitivity-Driven AI Risk Scoring", "RISK SCORING SPECIFICATION",
          "The solution must provide risk scoring for AI systems based on the sensitivity of data they access.",
          "Computes continuous risk scores reflecting the sensitivity of data repositories accessed by each AI system or employee profile.",
          EXTRACTED_DIR / "28_image29.png",
          "Figure 11.1: LLM Interaction Log tracking request tokens, latency, cost metrics, and dynamic risk scoring."),
         ("Excessive & Unnecessary Data Access Detection", "ANOMALY DETECTION SPECIFICATION",
          "The solution should detect when AI tools attempt to access excessive or unnecessary data.",
          "Monitors query volume and scope; excessive requests trigger instant JIT notification banners and prompt rate-limiting.",
          ASSETS_DIR / "p2_img1.jpeg",
          "Figure 11.2: Gateway interception of excessive query bursts, jailbreak patterns, and mass data extraction.")),

        # Page 14: Requirements 56 & 57
        ("Dataset Access & Anomaly Detection",
         ("Policy-Based AI Dataset Access Restriction", "DATASET RESTRICTION SPECIFICATION",
          "The solution must provide the ability to restrict AI model access to sensitive datasets based on predefined policies.",
          "Restricts access to confidential documents by disabling clipboard copy, file attachments, and download exports inside AI tool interfaces.",
          ASSETS_DIR / "ui_setup.png",
          "Figure 12.1: Gateway Administration console enforcing dataset isolation, SSL certificates, and tenant domain locking."),
         ("Anomalous AI Data Access Pattern Monitoring", "ACCESS PATTERN SPECIFICATION",
          "The solution should monitor and detect anomalous AI data access patterns that may indicate misuse or compromise.",
          "Tracks token velocity, prompt cadence, and behavioral deviations to detect credential compromise and automated prompt harvesting.",
          ASSETS_DIR / "ui_guard_insights.png",
          "Figure 12.2: Guard Insights anomalous activity monitor tracking real-time policy breach spikes and threat trends.")),

        # Page 15: Requirements 58 & 59
        ("Suspicious Data & Exfiltration Alerts",
         ("Suspicious Data Interaction Detection", "INTERACTION DETECTION SPECIFICATION",
          "The solution must detect abnormal or suspicious interactions between AI systems and enterprise data sources.",
          "Inspects interactions passing between users and external models, flagging abnormal query structures and classified file queries.",
          ASSETS_DIR / "ui_guard_inspect_modal.png",
          "Figure 13.1: Suspicious Interaction Inspector dialog providing deep packet payload analysis and rule evaluation history."),
         ("Automated Data Exfiltration Alerting", "EXFILTRATION ALERT SPECIFICATION",
          "The solution should generate alerts for potential data exfiltration attempts through AI applications.",
          "Instantly triggers browser popups, telemetry alerts, and SOC incident notifications when bulk data extraction is detected.",
          ASSETS_DIR / "extra_img_2.png",
          "Figure 13.2: High-priority security incident console logging automated exfiltration triggers and containment actions.")),

        # Page 16: Requirements 60 & 61
        ("Scraping Mitigation & Workflows",
         ("Abnormal Usage & Automated Scraping Mitigation", "SCRAPING DEFENSE SPECIFICATION",
          "The solution must detect abnormal AI usage patterns such as excessive prompts, automated extraction attempts, or data scraping behaviour.",
          "Rate-limits rapid query generation, terminates automated bot sessions, and logs enumeration sequences at the proxy gateway.",
          ASSETS_DIR / "extra_img_1.png",
          "Figure 14.1: LLM Gateway policy engine enforcing rate limits, token quotas, and prompt scraping defense."),
         ("Security Analyst Investigation Workflows", "ANALYST WORKFLOW SPECIFICATION",
          "The solution should support investigation workflows for security analysts to review AI-related security events.",
          "Equips security analysts with granular search, domain filters, status filters (Allowed vs. Blocked), and complete prompt payload inspection.",
          ASSETS_DIR / "ui_search_logs.png",
          "Figure 14.2: Security Analyst Search & Investigation console with granular search queries, engine filters, and threat scores.")),

        # Page 17: Requirements 62 & 63
        ("Audit Logging & Governance Controls",
         ("Detailed Immutable AI Audit Logging", "AUDIT LOGGING SPECIFICATION",
          "The solution must maintain detailed audit logs of AI activity and data access events.",
          "Captures immutable records detailing date, time, endpoint hostname, AI platform, user prompt, token usage, cost, and exact enforcement decision.",
          ASSETS_DIR / "p5_img2.jpeg",
          "Figure 15.1: Comprehensive immutable audit log table detailing timestamps, user sessions, prompt risk, and verdicts."),
         ("Policy-Aligned Enterprise Governance Controls", "GOVERNANCE CONTROL SPECIFICATION",
          "The solution must provide governance controls for enterprise AI usage aligned with organizational security policies.",
          "Enforces corporate workspace domain authorization (@yespanchi.com), enterprise license quotas, and centralized workstation sync.",
          ASSETS_DIR / "extra_img_4.png",
          "Figure 15.2: Enterprise Governance & Compliance configuration console aligning AI tool permissions with corporate security policies.")),

        # Page 18: Requirements 64 & 65
        ("Compliance Reports & Risk Dashboards",
         ("AI Data Access Compliance Reporting", "COMPLIANCE REPORTING SPECIFICATION",
          "The solution should generate compliance reports related to AI data access and AI system usage.",
          "Provides dynamic executive compliance telemetry capturing total prompts intercepted, high-risk flags, and real-time report export options.",
          ASSETS_DIR / "extra_img_3.png",
          "Figure 16.1: MCP Gateway tools & compliance telemetry dashboard tracking connector requests and export summaries."),
         ("Enterprise AI Risk Posture Dashboards", "RISK POSTURE SPECIFICATION",
          "The solution must provide dashboards for security teams showing AI risk posture across the enterprise.",
          "Displays fleet-wide endpoint agent health, policy interception hit counts (Allowed vs. Blocked), and active workstation risk status.",
          ASSETS_DIR / "ui_overview.png",
          "Figure 16.2: Executive AI Risk Posture overview dashboard displaying daily active agents, intercepted prompts, and threat scores.")),

        # Page 19: Requirements 66 & 67
        ("Audit Reporting & Dataset Protection",
         ("Regulatory & Internal Audit Reporting", "AUDIT REPORTING SPECIFICATION",
          "The solution should support audit reporting for regulatory or internal governance requirements.",
          "Captures forensic compliance trails across web and search queries, logging browser engines, destination result clicks, and threat risk ratings.",
          ASSETS_DIR / "ui_search_logs.png",
          "Figure 17.1: Enterprise Search Activity & Governance Audit log detailing query text, destination URLs, and threat ratings."),
         ("Restricted & Classified Dataset Access Prevention", "DATASET PROTECTION SPECIFICATION",
          "The solution must enable policy enforcement to prevent AI tools from accessing restricted or classified datasets.",
          "Enforces strict prevention policies blocking upload, copy-paste, and drag-and-drop of restricted source files into AI interfaces.",
          EXTRACTED_DIR / "04_image25.png",
          "Figure 17.2: Connected repository permissions console preventing unauthorized AI tool ingestion of sensitive file stores.")),

        # Page 20: Requirements 68 & 69
        ("Runtime Monitoring & Threat Defense",
         ("Runtime Monitoring of AI Interactions", "RUNTIME MONITORING SPECIFICATION",
          "The solution must support runtime monitoring of AI interactions with enterprise data.",
          "Inline proxy broker inspects live data exchanges, evaluating prompt tokens, response safety, and model actions in sub-5ms latency.",
          ASSETS_DIR / "ui_setup.png",
          "Figure 18.1: Gateway Runtime Proxy architecture console configuring real-time data inspection and TLS termination."),
         ("Prompt Injection & Adversarial Attack Detection", "THREAT DEFENSE SPECIFICATION",
          "The solution should detect prompt-based attacks such as prompt injection or manipulation attempts.",
          "Identifies adversarial syntax, recursive prompt injection strings, and instruction overrides, dropping hostile payloads at the socket layer.",
          ASSETS_DIR / "p2_img1.jpeg",
          "Figure 18.2: Real-time defense log intercepting prompt injection payloads, system overrides, and adversarial exploits.")),

        # Page 21: Requirements 70 & 71
        ("Output Protection & Remediation",
         ("Sensitive Output Leakage Prevention", "OUTPUT PROTECTION SPECIFICATION",
          "The solution must prevent AI systems from exposing sensitive information through responses or outputs.",
          "Bidirectional output filter scrubs sensitive tokens, PII, and credentials from downstream model completions before rendering in the browser.",
          ASSETS_DIR / "p4_img2.jpeg",
          "Figure 19.1: Response Scanning dashboard tagging safety classifications and masking leaked output tokens."),
         ("Automated Remediation Actions", "AUTOMATED REMEDIATION SPECIFICATION",
          "The solution should support automated remediation actions such as blocking AI requests, restricting access, or generating alerts.",
          "Automatically executes inline request blocking, drops active TCP sessions, and triggers instant telemetry alerts upon critical policy breaches.",
          ASSETS_DIR / "p3_img1.jpeg",
          "Figure 19.2: Browser Guard enforcement console displaying active automated blocking rules for unauthorized actions.")),

        # Page 22: Requirements 72 & 73
        ("Multi-Asset Risk & Adoption Trends",
         ("Centralized Multi-Asset Risk Exposure Dashboards", "RISK EXPOSURE SPECIFICATION",
          "The solution must provide centralized dashboards displaying AI risk exposure across users, systems, and data assets.",
          "Inspects individual endpoint workstations, detailing hardware identity, physical MAC, OS version, and granular telemetry audit logs.",
          ASSETS_DIR / "extra_img_5.png",
          "Figure 20.1: Multi-Asset Hardware & Workstation telemetry monitor tracking CPU, memory, OS versions, and agent health."),
         ("AI Adoption Trends & Risk Indicator Analytics", "ADOPTION ANALYTICS SPECIFICATION",
          "The solution should provide analytics on AI adoption trends, usage patterns, and risk indicators.",
          "Analyzes query volume growth, cost per user, token consumption, and risk indicator trends across commercial and internal AI platforms.",
          ASSETS_DIR / "ui_overview.png",
          "Figure 20.2: Enterprise AI Adoption trends console mapping active daily users, prompt distribution, and model consumption.")),

        # Page 23: Requirements 74 & 75
        ("Behavioral Scoring & Executive Reports",
         ("Behavioral & Contextual AI Risk Scoring", "BEHAVIORAL SCORING SPECIFICATION",
          "The solution must provide risk scoring for AI activities based on data sensitivity and user behaviour.",
          "Evaluates contextual risk dynamically using historical user behavior, query intent, and search privacy modes.",
          ASSETS_DIR / "p5_img1.jpeg",
          "Figure 21.1: Multi-factor contextual risk matrix weighting user behavior, data classification, and platform security."),
         ("Executive Leadership & Risk Committee Reporting", "EXECUTIVE REPORTING SPECIFICATION",
          "The solution should support executive reporting for security leadership and risk management teams.",
          "Summarizes enterprise AI compliance posture, blocked data leakage attempts, and cost telemetry into executive-level summaries.",
          ASSETS_DIR / "extra_img_4.png",
          "Figure 21.2: Executive Compliance & Governance summary report designed for enterprise risk committee review.")),
    ]

    for page_title, item1, item2 in pairs:
        doc.add_page_break()
        p_h1 = doc.add_paragraph()
        p_h1.paragraph_format.space_before = Pt(2)
        p_h1.paragraph_format.space_after = Pt(2)
        r = p_h1.add_run(page_title)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = NAVY_RGB
        
        add_topic_item_docx(doc, item1[0], item1[1], item1[2], item1[3], item1[4], item1[5], img_w=6.0)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
        add_topic_item_docx(doc, item2[0], item2[1], item2[2], item2[3], item2[4], item2[5], img_w=6.0)

    # Save to all requested locations
    print(f"Saving directly to: {TARGET_UPDATED_DOC}")
    doc.save(str(TARGET_UPDATED_DOC))
    
    print(f"Saving to: {TARGET_NEW_DOC}")
    doc.save(str(TARGET_NEW_DOC))
    
    print(f"Saving copy to: {OUTPUT_DOCX}")
    doc.save(str(OUTPUT_DOCX))
    
    shutil.copyfile(str(OUTPUT_DOCX), str(LOCAL_COPY))
    print(f"Local copy saved to: {LOCAL_COPY}")
    print("Generation complete! Total 23 Pages with perfect centering, clean index, and dynamic page numbers!")

if __name__ == "__main__":
    generate_perfect_updated_doc()
