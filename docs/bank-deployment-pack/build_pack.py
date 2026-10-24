from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent
LOGO = ROOT.parents[1] / "ui" / "public" / "yes-panchi-logo.png"
NAVY = "18304F"
BLUE = "2A638C"
PALE = "EAF1F7"
PALE2 = "F6F8FA"
GRID = "D9D9D9"
TEXT = "1D2733"
MUTED = "596879"
DOCS = [
    {
        "source": "01-technical-documentation.md",
        "stem": "01-yp-raksha-technical-documentation",
        "title": "YP-Raksha Technical Documentation",
        "subtitle": "Architecture, APIs, data model, security controls and Browser AI implementation",
        "kind": "TECHNICAL MANUAL",
    },
    {
        "source": "02-product-datasheet.md",
        "stem": "02-yp-raksha-product-datasheet",
        "title": "YP-Raksha Gateway and Browser AI Guard Product Datasheet",
        "subtitle": "Capabilities, interfaces, platform support, data handling and prerequisites",
        "kind": "PRODUCT REFERENCE",
    },
    {
        "source": "03-on-premises-deployment-architecture.md",
        "stem": "03-on-premises-deployment-architecture",
        "title": "YP-Raksha On-Premises Deployment Architecture",
        "subtitle": "Reference design for 120,000 endpoints across DC, DR and UAT",
        "kind": "ON-PREMISES ARCHITECTURE",
    },
]

FLOW_DIAGRAMS = {
    "flow-deployment": {
        "file": "deployment-sequence.png",
        "title": "On-premises release and endpoint rollout sequence",
        "steps": [
            ("Artifact intake", ["Verify signature and digest", "Scan and mirror internally"]),
            ("Platform and data", ["Provision DB and storage", "Separate DC, DR and UAT"]),
            ("Gateway deployment", ["Apply protected configuration", "Start and check readiness"]),
            ("Internal services", ["Private model and identity", "Internal logs and telemetry"]),
            ("Endpoint agent", ["Deploy signed Windows/macOS", "Confirm healthy heartbeat"]),
            ("Browser extension", ["Managed Chromium / Firefox", "Verify agent hand-off"]),
            ("UAT and promotion", ["Test no-egress and load", "Promote identical release"]),
        ],
        "note": "Site addresses, listener values, database identifiers and credentials stay in the bank's protected configuration records.",
    },
    "flow-inference": {
        "file": "inference-flow.png",
        "title": "Application inference request and response",
        "steps": [
            ("Client request", ["HTTPS API call", "Application or SDK"]),
            ("Internal ingress", ["TLS, WAF and limits", "WebSocket route if used"]),
            ("Authenticate", ["Virtual key or session", "Identity and scope"]),
            ("Validate + govern", ["Request and model", "RBAC, budget, policy"]),
            ("Select route", ["Rule and provider", "Approved key / target"]),
            ("Provider adapter", ["Private model service", "Bounded retry / fallback"]),
            ("Return response", ["Normalize or stream", "HTTP / event semantics"]),
            ("Record outcome", ["Configured LogsStore", "Internal metrics / trace"]),
        ],
        "note": "Model output returns through the Gateway to the client. Optional logs, traces and providers require approved internal destinations.",
    },
    "flow-admin-ui": {
        "file": "admin-ui-flow.png",
        "title": "Administration UI and policy change",
        "steps": [
            ("Administrator", ["Internal browser", "Named account"]),
            ("Embedded UI", ["Served with Gateway", "Internal ingress"]),
            ("Session check", ["Authentication", "Session / token"]),
            ("Authorize action", ["Resource + operation", "Section grants"]),
            ("Validate change", ["Feature handler", "Payload and state"]),
            ("Persist config", ["ConfigStore", "Approved DB service"]),
            ("Audit change", ["Actor and operation", "Internal audit records"]),
            ("Verify rollout", ["Agent/config pull", "UI + endpoint check"]),
        ],
        "note": "Use least privilege and confirm the saved object, audit record and endpoint effect before closing a policy change.",
    },
    "flow-identity-access": {
        "file": "identity-access-flow.png",
        "title": "Bank identity, session and authorization lifecycle",
        "steps": [
            ("Provision identity", ["Local/admin path or SCIM*", "Bank-controlled lifecycle"]),
            ("Open workspace", ["Managed employee browser", "Internal TLS ingress"]),
            ("Authenticate", ["Configured login/session", "OAuth2 where enabled"]),
            ("Establish session", ["Cookie or bearer session", "According to API surface"]),
            ("Authorize route", ["Resource × operation", "Permission check"]),
            ("Limit visible scope", ["Section grants", "Unknown role = deny"]),
            ("Execute or refuse", ["Permitted UI/API action", "401/403 if access fails"]),
            ("Audit mutation", ["Mutating management action", "Internal audit records"]),
        ],
        "note": "*SCIM is optional provisioning, not session authentication. OAuth2/SSO behavior depends on the selected release and configuration; bank MFA remains an identity-control responsibility.",
    },
    "flow-workspace-lifecycle": {
        "file": "workspace-lifecycle-flow.png",
        "title": "Prompt and Skills workspace lifecycle",
        "steps": [
            ("Open repository", ["Prompt or Skills page", "Authenticated session"]),
            ("Check access", ["Role and operation rights", "Section/repository scope"]),
            ("Browse content", ["Visible folders and items", "Authorized scope only"]),
            ("Create or edit", ["Prompt or skill content", "Permitted actions only"]),
            ("Save record", ["Repository data in bank DB", "Supported metadata/state"]),
            ("Review history", ["Versions where supported", "Inspect prior content"]),
            ("Deploy or serve", ["Prompt deployment API", "Skill serving if enabled"]),
            ("Retain and audit", ["Mutation audit as applicable", "Bank retention/backup rules"]),
        ],
        "note": "Repository features, version/history behavior and deployment or serving paths depend on enabled modules and release. Treat saved content as bank data.",
    },
    "flow-mcp": {
        "file": "mcp-flow.png",
        "title": "Optional MCP tool request",
        "steps": [
            ("Approved client", ["Tool request", "User / app identity"]),
            ("Gateway checks", ["Session and access", "Tool group policy"]),
            ("MCP connection", ["Configured transport", "Approved credentials"]),
            ("Bank MCP server", ["Private destination", "Allowed tool executes"]),
            ("Tool result", ["Validated response", "Error mapped if needed"]),
            ("Return + record", ["Client response", "MCP logs if enabled"]),
        ],
        "note": "Every MCP server is a data destination. Use bank-operated servers and an explicit host/port allow-list for no-egress deployment.",
    },
    "flow-release": {
        "file": "release-promotion-flow.png",
        "title": "Controlled build and promotion pipeline",
        "steps": [
            ("Source and review", ["Approved source baseline", "No source build on prod nodes"]),
            ("Build and package", ["Pinned toolchain and inputs", "Versioned image, chart, Guard"]),
            ("Automated tests", ["Go, Python and UI tests", "Synthetic integration checks"]),
            ("Security review", ["Scan dependencies/artifacts", "Review configuration and SBOM"]),
            ("Sign and identify", ["Approve signed release", "Record immutable digests"]),
            ("Internal mirror", ["Controlled artifact import", "No public runtime pulls"]),
            ("UAT qualification", ["Migration and policy checks", "Load, DR and no-egress tests"]),
            ("Promote and recover", ["Same digest to DR and DC", "Approved rollback/restore path"]),
        ],
        "note": "The bank controls signing trust, release approval, internal repository access, environment secrets and production promotion.",
    },
    "flow-observability": {
        "file": "observability-settings-flow.png",
        "title": "Settings, runtime signals and observability lifecycle",
        "steps": [
            ("Admin change", ["Internal workspace", "Named account"]),
            ("Authorize", ["Session + RBAC", "Audit mutation"]),
            ("Validate + save", ["Feature API", "ConfigStore"]),
            ("Apply settings", ["Gateway runtime", "Guard policy polling"]),
            ("Serve requests", ["Inference / Guard", "Health and outcomes"]),
            ("Record signals", ["Configured logs", "Metrics / traces"]),
            ("Internal collect", ["Bank collectors", "Restricted access"]),
            ("Review + alert", ["Internal dashboards", "Approved response"]),
        ],
        "note": "Keep collectors, alert channels and exported diagnostics inside the approved bank/India boundary; payload logging is separately configured.",
    },
    "flow-error-handling": {
        "file": "error-recovery-flow.png",
        "title": "Gateway error classification and caller recovery",
        "steps": [
            ("Receive request", ["Assign / retain ID", "Apply deadline"]),
            ("Validate access", ["Auth, schema, policy", "4xx: correct request"]),
            ("Call approved route", ["Private provider", "Capture attempt result"]),
            ("Classify failure", ["Client vs dependency", "Consult retry policy"]),
            ("Bounded retry", ["Eligible 5xx / network", "Same key for transport"]),
            ("Rotate if eligible", ["Provider auth / quota", "Only with key pool"]),
            ("Return outcome", ["HTTP before stream", "Protocol error after start"]),
            ("Reconcile safely", ["No blind side-effect replay", "Use ID / provider state"]),
        ],
        "note": "Actual retry classes and limits are release/configuration dependent. Do not replay a non-idempotent request just because the caller timed out.",
    },
    "flow-guard-rollout": {
        "file": "onprem-guard-rollout.png",
        "title": "Bank-managed Guard agent and browser-extension rollout",
        "steps": [
            ("Approve release", ["Verify package digest", "Confirm OS/browser scope"]),
            ("Prepare endpoint", ["Internal Gateway name", "Protected fleet settings"]),
            ("Install agent", ["Windows or macOS package", "Trust and startup behavior"]),
            ("Apply browser policy", ["Managed extension package", "Internal update source"]),
            ("Connect and verify", ["PAC, proxy and certificate", "Heartbeat and event path"]),
            ("Run synthetic cases", ["Allow, redact, warn, block", "Target and upload behavior"]),
            ("Expand in cohorts", ["Measure coverage and errors", "Rollback on regressions"]),
        ],
        "note": "The desktop agent and browser extension are separate components. Exact package names, signing, policy IDs and versions are release-specific.",
    },
    "flow-site-failover": {
        "file": "onprem-site-failover.png",
        "title": "Controlled DC-to-DR failover and failback",
        "steps": [
            ("Detect site incident", ["Review service and WAN state", "Open bank incident control"]),
            ("Freeze changes", ["Stop config promotion", "Protect current evidence"]),
            ("Assess recovery point", ["Measure DB lag and backup", "Confirm keys and certificates"]),
            ("Authorize promotion", ["DBA selects safe recovery", "Record measured RPO"]),
            ("Activate DR services", ["Promote approved database", "Check private dependencies"]),
            ("Scale and route", ["Scale Gateway capacity", "Move approved internal route"]),
            ("Validate service", ["Model, Guard and admin paths", "Observe errors and lag"]),
            ("Reconcile and fail back", ["Synchronize and compare data", "Controlled route cutback"]),
        ],
        "note": "The product does not automatically promote the bank database or perform cross-site failover. Database fencing and routing remain bank responsibilities.",
    },
    "flow-backup-restore": {
        "file": "onprem-backup-restore.png",
        "title": "Coordinated bank-resident backup and restore",
        "steps": [
            ("Select recovery point", ["Database plus attachments", "Capture config and release identity"]),
            ("Protect key material", ["Use bank key custody", "Restrict recovery operators"]),
            ("Encrypt and copy", ["Store only in India sites", "Record job and integrity result"]),
            ("Restore in isolation", ["Restore database and files", "Apply matched secrets and release"]),
            ("Reconcile records", ["Check file metadata and bytes", "Verify identity and retention"]),
            ("Resume approved route", ["Run synthetic service checks", "Record restore outcome"]),
        ],
        "note": "Database metadata, attachment bytes, configuration and recovery keys have different lifecycles; restore and verify the coordinated set.",
    },
    "flow-db-migration": {
        "file": "onprem-db-migration.png",
        "title": "Release-matched database migration path",
        "steps": [
            ("Import release", ["Verify image and chart", "Read migration notes"]),
            ("Recover UAT copy", ["Use masked or synthetic data", "Keep prod writes isolated"]),
            ("Test migration", ["Use controlled DB identity", "Measure locks and duration"]),
            ("Verify application", ["Check reads and writes", "Exercise Guard and inference"]),
            ("Protect production", ["Take coordinated recovery point", "Freeze conflicting changes"]),
            ("Promote identical release", ["Apply approved manifest", "Observe DB and app health"]),
            ("Recover if needed", ["Use tested restore path", "Do not assume downgrade safety"]),
        ],
        "note": "Database schema and migration behavior are release-specific. DBA and platform teams own HA, backup, promotion and restore procedures.",
    },
}


def set_font(run, name="Calibri", size=10.5, bold=None, color=TEXT, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), GRID)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def clean_inline(text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = text.replace("*", "").replace("_", "")
    return text.strip()


def clean_heading(text):
    text = re.sub(r"^\s*\d+(?:\.\d+)*\.?\s+", "", text.strip())
    text = re.sub(r"\s*\([^)]*/[^)]*\)", "", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"[^A-Za-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def add_inline(paragraph, text, base_size=10.5, base_color=TEXT):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    pattern = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+\*)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_font(run, size=base_size, color=base_color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=base_size, bold=True, color=base_color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_font(run, name="Consolas", size=max(8, base_size - 1), color="34495E")
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=base_size, italic=True, color=base_color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_font(run, size=base_size, color=base_color)


def normalize(text):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", text.casefold())).strip()


def font_path(bold=False):
    candidates = [
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    return next((p for p in candidates if p.exists()), candidates[-1])


def draw_architecture(path, product=False):
    width, height = ((1900, 1020) if product else (1900, 1120))
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f_title = ImageFont.truetype(str(font_path(True)), 38)
    f_label = ImageFont.truetype(str(font_path(True)), 25)
    f_body = ImageFont.truetype(str(font_path()), 21)
    f_small = ImageFont.truetype(str(font_path()), 18)

    def box(xy, title, lines, fill="#FFFFFF", outline="#35688F", radius=16):
        draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=3)
        x0, y0, x1, _ = xy
        draw.text((x0 + 18, y0 + 15), title, font=f_label, fill="#18304F")
        y = y0 + 56
        for line in lines:
            draw.text((x0 + 18, y), line, font=f_body, fill="#293744")
            y += 33

    def arrow(start, end, color="#2A638C", width_px=5, dashed=False, label=None, label_xy=None):
        x1, y1 = start
        x2, y2 = end
        if dashed:
            dx, dy = x2 - x1, y2 - y1
            length = max(1, (dx * dx + dy * dy) ** 0.5)
            ux, uy = dx / length, dy / length
            step = 18
            pos = 0
            while pos < length - 24:
                endpos = min(pos + 12, length - 24)
                draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * endpos, y1 + uy * endpos), fill=color, width=width_px)
                pos += step
        else:
            draw.line((x1, y1, x2, y2), fill=color, width=width_px)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 18
        left = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
        right = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), left, right], fill=color)
        if label:
            lx, ly = label_xy if label_xy else ((x1 + x2) // 2, (y1 + y2) // 2 - 24)
            draw.text((lx, ly), label, font=f_small, fill=color)

    if product:
        draw.text((56, 30), "Logical product architecture", font=f_title, fill="#182635")
        draw.rounded_rectangle((35, 90, 1865, 950), radius=24, fill="#FBFCFD", outline="#18304F", width=4)
        draw.text((62, 103), "BANK-CONTROLLED TRUST BOUNDARY  |  deploy approved components and dependencies inside this zone", font=f_label, fill="#18304F")
        box((75, 215, 365, 445), "Consumers", ["Internal AI applications", "Platform administrators", "Managed browser users"], fill="#F2F6FA")
        box((75, 620, 365, 830), "Browser AI Guard", ["Endpoint agent + local proxy", "Optional internal shared proxy", "PAC, policy and event exchange"], fill="#F5F1F7", outline="#775A87")
        box((435, 265, 700, 515), "Internal ingress", ["WAF / load balancer", "TLS termination", "Pass-through TLS", "WebSocket-aware routing", "Network allow-list"], fill="#F7F5EF", outline="#8A7446")
        box((770, 175, 1265, 700), "YP-Raksha Gateway", ["Embedded administration UI + APIs", "Authentication, sessions and RBAC", "Governance, budgets and audit", "Policy, guardrails and model routing", "Provider / MCP adapters", "Health, metrics and request tracing"], fill="#EAF1F7")
        box((1380, 170, 1805, 365), "Configuration and identity", ["ConfigStore: users, roles, keys", "Providers, models and policy", "PostgreSQL reference baseline"], fill="#F7F5EF", outline="#8A7446")
        box((1380, 430, 1805, 655), "Logs and event records", ["Inference, audit and MCP logs", "Browser AI records + metadata", "Payloads only when configured"], fill="#F7F5EF", outline="#8A7446")
        box((1380, 715, 1805, 875), "Private operations services", ["Approved internal model endpoints", "Metrics / logs / trace collectors", "Backup, artifact and secret stores"], fill="#EFF6F1", outline="#4A8062")
        arrow((365, 330), (435, 365))
        arrow((700, 385), (770, 385))
        arrow((365, 710), (435, 465))
        arrow((1265, 285), (1380, 270))
        arrow((1265, 455), (1380, 520))
        arrow((1265, 600), (1380, 785))
        draw.text((75, 905), "Applications and Guard use internal ingress; Gateway reaches only approved private models, stores and operations services.", font=f_small, fill="#536273")
    else:
        draw.text((55, 25), "Sovereign deployment data flows across DC, DR and UAT", font=f_title, fill="#182635")
        draw.rounded_rectangle((45, 85, 1855, 845), radius=26, fill="#F6F9FC", outline="#18304F", width=5)
        draw.text((70, 98), "BANK-CONTROLLED NETWORKS IN INDIA  |  production and test data remain inside this boundary", font=f_label, fill="#18304F")
        box((80, 220, 360, 450), "User zones", ["Managed browser + Guard", "Internal AI applications", "Admin workstations"], fill="#F2F6FA")
        box((445, 180, 1050, 670), "Primary DC", ["Internal WAF / ingress", "Gateway replicas", "Private model endpoint", "PostgreSQL ConfigStore + LogsStore", "Internal metrics, logs and backups"], fill="#EAF1F7")
        box((1135, 180, 1450, 670), "DR site", ["Warm Gateway capacity", "Replicated database", "Internal collectors", "Encrypted bank WAN"], fill="#EFF6F1", outline="#4A8062")
        box((1510, 180, 1815, 670), "Isolated UAT", ["Separate Gateway and DB", "Synthetic / masked data", "Private test model", "No production DB access"], fill="#F7F5EF", outline="#8A7446")
        arrow((360, 320), (445, 320))
        arrow((1050, 390), (1135, 390))
        draw.text((70, 735), "Gateway → private models   •   logs/metrics → internal collectors   •   encrypted backups/replicas → bank DR site in India", font=f_small, fill="#293744")
        box((80, 930, 570, 1080), "Public AI and model APIs", ["DENY — no prompt or response egress"], fill="#FFF4F2", outline="#A13A36", radius=12)
        box((675, 930, 1165, 1080), "External log and telemetry services", ["DENY — no logs, traces or metrics"], fill="#FFF4F2", outline="#A13A36", radius=12)
        box((1270, 930, 1815, 1080), "External registry, licensing and updates", ["DENY — use internal mirrors / approved offline flow"], fill="#FFF4F2", outline="#A13A36", radius=12)
        for x in (325, 920, 1540):
            arrow((x, 845), (x, 925), color="#A13A36", width_px=4, dashed=True)
            draw.line((x - 13, 870, x + 13, 896), fill="#A13A36", width=5)
            draw.line((x + 13, 870, x - 13, 896), fill="#A13A36", width=5)
    image.save(path, quality=95)


def draw_platform_architecture(path):
    """Draw the full YP-Raksha client, control, data, integration and operations architecture."""
    width, height = 1900, 1130
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f_title = ImageFont.truetype(str(font_path(True)), 38)
    f_label = ImageFont.truetype(str(font_path(True)), 25)
    f_body = ImageFont.truetype(str(font_path()), 21)
    f_note = ImageFont.truetype(str(font_path()), 19)
    navy, blue, muted = "#18304F", "#2A638C", "#536273"

    draw.text((48, 23), "YP-Raksha — full product and integration architecture", font=f_title, fill="#182635")
    draw.rounded_rectangle((28, 83, 1872, 1080), radius=24, fill="#FBFCFD", outline=navy, width=4)
    draw.text((52, 94), "BANK / INDIA TRUST BOUNDARY  |  bank-operated services, stores and approved private destinations", font=f_label, fill=navy)

    def box(x, y, w, h, title, lines, fill, outline="#35688F"):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=16, fill=fill, outline=outline, width=3)
        draw.text((x + 15, y + 13), title, font=f_label, fill=navy)
        max_width = w - 30
        cursor_y = y + 51
        for item in lines:
            words = item.split()
            current = ""
            for word in words:
                candidate = f"{current} {word}".strip()
                if current and draw.textbbox((0, 0), candidate, font=f_body)[2] > max_width:
                    draw.text((x + 15, cursor_y), current, font=f_body, fill="#293744")
                    cursor_y += 29
                    current = word
                else:
                    current = candidate
            if current:
                draw.text((x + 15, cursor_y), current, font=f_body, fill="#293744")
                cursor_y += 29

    def arrow(start, end, color=blue, dashed=False, width_px=4):
        x1, y1 = start
        x2, y2 = end
        if dashed:
            dx, dy = x2 - x1, y2 - y1
            length = max(1, (dx * dx + dy * dy) ** 0.5)
            ux, uy = dx / length, dy / length
            pos = 0
            while pos < length - 20:
                endpos = min(pos + 11, length - 20)
                draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * endpos, y1 + uy * endpos), fill=color, width=width_px)
                pos += 18
        else:
            draw.line((x1, y1, x2, y2), fill=color, width=width_px)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 15
        left = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
        right = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), left, right], fill=color)

    # Control plane: users and API clients, access controls, configuration and Guard fleet policy.
    control_x = [55, 425, 795, 1165, 1535]
    control_titles = [
        ("Workspace access", ["Employees, app clients, admins", "Named accounts and service identities"]),
        ("YP-Raksha UI + APIs", ["Admin workspace and management", "Inference and feature APIs"]),
        ("Identity + authorization", ["Session / API authentication", "RBAC, section grants, audit", "Optional directory / SCIM flow"]),
        ("Policy + configuration", ["Providers, models and virtual keys", "Budgets, routes, Guard rules", "Targets, prompts, skills, plugins"]),
        ("Guard fleet control", ["Enrollment and policy delivery", "Heartbeat, polling and event path"]),
    ]
    for x, (title, lines) in zip(control_x, control_titles):
        box(x, 138, 330, 178, title, lines, "#F2F6FA" if x != control_x[3] else "#F7F5EF", "#8A7446" if x == control_x[3] else "#35688F")
    for i in range(4):
        arrow((control_x[i] + 330, 227), (control_x[i + 1] - 7, 227))
    draw.text((65, 324), "CONTROL PLANE  •  authorized settings are validated, persisted and consumed by the relevant runtime", font=f_note, fill=muted)

    # Data plane: app and Guard traffic through the Gateway to approved private services.
    data_x = [55, 360, 665, 970, 1275, 1580]
    data_items = [
        ("Client + Guard", ["Internal apps / API clients", "Managed browser extension", "Endpoint agent / local proxy"], "#F2F6FA"),
        ("Bank ingress", ["Internal WAF / load balancer", "TLS and route controls", "Streaming / WebSocket as needed"], "#F7F5EF"),
        ("YP-Raksha Gateway", ["Inference and management APIs", "Guard policy / event APIs", "Prompt, Skills and optional MCP APIs"], "#EAF1F7"),
        ("Policy + routing", ["Authentication, keys, budgets", "Guard rules, targets, verdicts", "Routes, bounded retry / breakers"], "#EAF1F7"),
        ("Private adapters", ["Configured model adapters", "Optional MCP / cache / vector", "Policy applies per route"], "#EFF6F1"),
        ("Private destinations", ["Private model endpoints", "Bank MCP (optional)", "Cache/vector in bank", "Public denied in strict mode"], "#EFF6F1"),
    ]
    for x, (title, lines, fill) in zip(data_x, data_items):
        box(x, 405, 270, 235, title, lines, fill, "#4A8062" if fill == "#EFF6F1" else "#35688F")
    for i in range(5):
        arrow((data_x[i] + 270, 522), (data_x[i + 1] - 7, 522))
    # Persistence and observability plane.
    store_x = [55, 425, 795, 1165, 1535]
    store_items = [
        ("Configuration data", ["Users, roles and settings", "Provider/model configuration", "Policy, prompts and governance"], "#F7F5EF"),
        ("Inference / MCP records", ["Configured request records", "Audit and tool activity", "Retention by record class"], "#F7F5EF"),
        ("Browser AI records", ["browser_ai_logs", "browser_ai_search_logs", "Attachment bytes on durable path"], "#F7F5EF"),
        ("Runtime signals", ["Health and service logs", "Internal metrics / traces", "Sensitive labels minimized"], "#EAF1F7"),
        ("Internal operations", ["Bank collectors and alerts", "Restricted dashboard access", "Encrypted backup / DR copies"], "#EFF6F1"),
    ]
    for x, (title, lines, fill) in zip(store_x, store_items):
        box(x, 695, 330, 190, title, lines, fill, "#4A8062" if fill == "#EFF6F1" else "#8A7446" if fill == "#F7F5EF" else "#35688F")
    arrow((800, 640), (590, 687), color="#6C7785", dashed=True)
    arrow((865, 640), (960, 687), color="#6C7785", dashed=True)
    arrow((1100, 640), (1330, 687), color="#6C7785", dashed=True)
    draw.text((65, 894), "PERSISTENCE + OBSERVABILITY  •  separate store classes; response returns on request path; capture and retention follow bank policy", font=f_note, fill=muted)

    draw.rounded_rectangle((55, 930, 1845, 1040), radius=13, fill="#F6F8FA", outline="#CBD5E1", width=2)
    draw.text((75, 946), "BANK-OPERATED DEPENDENCIES", font=f_label, fill=navy)
    draw.text((75, 983), "Bank identity / secrets • DNS / PKI / time • signed artifact mirror • relational DB / durable files • private model • internal collectors • encrypted DC ↔ DR", font=f_note, fill="#293744")
    draw.text((55, 1052), "NO-EGRESS: block public AI/model, SaaS telemetry/analytics, public update and unapproved connector paths at endpoint and network boundaries.", font=f_note, fill="#9B3434")
    image.save(path, quality=95)


def draw_flow_diagram(path, title, steps, note):
    width, height = 1900, 1000
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f_title = ImageFont.truetype(str(font_path(True)), 38)
    f_label = ImageFont.truetype(str(font_path(True)), 25)
    f_body = ImageFont.truetype(str(font_path()), 21)
    f_small = ImageFont.truetype(str(font_path()), 18)
    draw.text((55, 28), title, font=f_title, fill="#182635")
    draw.rounded_rectangle((35, 95, 1865, 900), radius=24, fill="#FBFCFD", outline="#CBD5E1", width=3)
    columns = 4
    margin, gap = 58, 28
    card_w = (width - 2 * margin - (columns - 1) * gap) // columns
    card_h, row_gap = 255, 105
    coords = []
    colors = ["#F2F6FA", "#EAF1F7", "#F7F5EF", "#EFF6F1"]
    for index, (label, lines) in enumerate(steps):
        row, offset = divmod(index, columns)
        col = offset if row % 2 == 0 else columns - 1 - offset
        x0 = margin + col * (card_w + gap)
        y0 = 145 + row * (card_h + row_gap)
        x1, y1 = x0 + card_w, y0 + card_h
        coords.append((x0, y0, x1, y1, row, col))
        draw.rounded_rectangle((x0, y0, x1, y1), radius=18, fill=colors[index % len(colors)], outline="#35688F", width=3)
        draw.text((x0 + 18, y0 + 18), f"{index + 1}. {label}", font=f_label, fill="#18304F")
        y = y0 + 70
        for line in lines:
            draw.text((x0 + 18, y), line, font=f_body, fill="#293744")
            y += 38

    def arrow(start, end, color="#2A638C", dashed=False):
        x1, y1 = start
        x2, y2 = end
        if dashed:
            dx, dy = x2 - x1, y2 - y1
            length = max(1, (dx * dx + dy * dy) ** 0.5)
            ux, uy = dx / length, dy / length
            pos = 0
            while pos < length - 18:
                endpos = min(pos + 11, length - 18)
                draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * endpos, y1 + uy * endpos), fill=color, width=4)
                pos += 18
        else:
            draw.line((x1, y1, x2, y2), fill=color, width=4)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 16
        left = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
        right = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), left, right], fill=color)

    for index in range(len(coords) - 1):
        x0, y0, x1, y1, row, col = coords[index]
        nx0, ny0, nx1, ny1, nrow, ncol = coords[index + 1]
        if row == nrow and ncol > col:
            arrow((x1, (y0 + y1) // 2), (nx0, (ny0 + ny1) // 2))
        elif row == nrow:
            arrow((x0, (y0 + y1) // 2), (nx1, (ny0 + ny1) // 2))
        else:
            start_x = (x0 + x1) // 2
            end_x = (nx0 + nx1) // 2
            arrow((start_x, y1), (end_x, ny0))
    draw.text((58, 925), note, font=f_small, fill="#536273")
    image.save(path, quality=95)


def draw_error_handling_diagram(path):
    """Draw conditional error branches without implying that retries always run in sequence."""
    width, height = 1900, 1030
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f_title = ImageFont.truetype(str(font_path(True)), 38)
    f_label = ImageFont.truetype(str(font_path(True)), 25)
    f_body = ImageFont.truetype(str(font_path()), 21)
    f_note = ImageFont.truetype(str(font_path()), 19)
    navy, blue = "#18304F", "#2A638C"
    draw.text((52, 27), "Gateway error classification and caller recovery", font=f_title, fill="#182635")
    draw.rounded_rectangle((32, 92, 1868, 950), radius=24, fill="#FBFCFD", outline="#CBD5E1", width=3)

    def box(x, y, w, h, title, lines, fill="#F2F6FA", outline="#35688F"):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=17, fill=fill, outline=outline, width=3)
        draw.text((x + 17, y + 16), title, font=f_label, fill=navy)
        text_y = y + 68
        for line in lines:
            words, current = line.split(), ""
            for word in words:
                candidate = f"{current} {word}".strip()
                if current and draw.textbbox((0, 0), candidate, font=f_body)[2] > w - 34:
                    draw.text((x + 17, text_y), current, font=f_body, fill="#293744")
                    text_y += 31
                    current = word
                else:
                    current = candidate
            if current:
                draw.text((x + 17, text_y), current, font=f_body, fill="#293744")
                text_y += 31

    def arrow(start, end, color=blue, width_px=4):
        import math
        x1, y1 = start
        x2, y2 = end
        draw.line((x1, y1, x2, y2), fill=color, width=width_px)
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 16
        left = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
        right = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), left, right], fill=color)

    top = [
        (55, 145, 375, 220, "1. Receive request", ["Retain request ID", "Apply configured deadline"]),
        (510, 145, 375, 220, "2. Authenticate + validate", ["Caller, session, schema", "Quota and policy checks"]),
        (965, 145, 375, 220, "3. Call approved route", ["Private provider endpoint", "Capture attempt outcome"]),
        (1420, 145, 425, 220, "4. Classify failure", ["Client, policy or dependency", "Provider, stream and retry class"]),
    ]
    for x, y, w, h, title, lines in top:
        box(x, y, w, h, title, lines, "#EAF1F7")
    for x1, x2 in [(430, 503), (885, 958), (1340, 1413)]:
        arrow((x1, 255), (x2, 255))

    # Branch from classification to the applicable failure path.
    draw.line((1632, 365, 1632, 414), fill="#6C7785", width=4)
    draw.line((345, 414, 1555, 414), fill="#6C7785", width=4)
    branches = [
        (85, "Client / auth / policy error", ["Return the relevant error", "Correct input, credential or grant", "Do not retry unchanged 4xx"], "#FFF4F2", "#A13A36"),
        (690, "Eligible network / 5xx / 429", ["Apply bounded retry policy", "Transport retry keeps same key", "Quota rotation only if eligible"], "#F7F5EF", "#8A7446"),
        (1295, "Provider credential / quota", ["Rotate only eligible provider key", "Retry within configured limits", "Stop when no eligible key remains"], "#EFF6F1", "#4A8062"),
    ]
    for x, title, lines, fill, outline in branches:
        center = x + 260
        draw.line((center, 414, center, 468), fill="#6C7785", width=4)
        arrow((center, 468), (center, 482), color="#6C7785")
        box(x, 485, 520, 205, title, lines, fill, outline)

    # Shared response rules apply after any selected branch.
    for center in (345, 950, 1555):
        draw.line((center, 690, center, 730), fill="#6C7785", width=3)
    draw.line((345, 730, 1555, 730), fill="#6C7785", width=3)
    arrow((950, 730), (950, 760), color="#6C7785")
    box(150, 765, 1600, 135, "Shared response and replay rules", [
        "Before output starts: return an HTTP status/error. After SSE/WebSocket output starts: end with protocol error or termination; the HTTP status cannot be replaced.",
        "For non-idempotent operations: do not blindly replay after timeout; reconcile with request ID or provider/tool state first.",
    ], "#F2F6FA", "#35688F")
    draw.text((55, 918), "Retry classes, limits and key-pool behavior depend on the selected release and configuration; provider failure does not imply automatic fallback.", font=f_note, fill="#536273")
    image.save(path, quality=95)


def draw_site_topology(path):
    """Draw the bank-specific site topology using the capacity values in the architecture reference."""
    width, height = 1900, 1330
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f_title = ImageFont.truetype(str(font_path(True)), 40)
    f_subtitle = ImageFont.truetype(str(font_path()), 26)
    f_label = ImageFont.truetype(str(font_path(True)), 29)
    f_body = ImageFont.truetype(str(font_path()), 28)
    f_note = ImageFont.truetype(str(font_path()), 22)

    draw.text((48, 24), "Reference deployment topology — DC, DR and UAT", font=f_title, fill="#182635")
    draw.text((50, 73), "Planning envelope for 120,000 managed endpoints; private model compute is sized separately.", font=f_subtitle, fill="#536273")
    draw.rounded_rectangle((28, 112, 1872, 1250), radius=25, fill="#FBFCFD", outline="#18304F", width=4)
    draw.text((52, 123), "BANK-CONTROLLED INDIA BOUNDARY", font=f_label, fill="#18304F")

    sites = [
        {
            "bounds": (48, 170, 650, 1125), "title": "PRIMARY DC — FULL PLANNED LOAD", "fill": "#EAF1F7", "outline": "#35688F",
            "items": [
                ("Ingress / WAF", ["2 × 8 vCPU / 16 GiB", "Active-active"]),
                ("Gateway worker pool", ["4 × 32 vCPU / 128 GiB nodes", "6 pods start → scale to 12"]),
                ("PostgreSQL HA", ["3 × 32 vCPU / 256 GiB nodes", "4 TiB usable; primary + replicas"]),
                ("Private dependencies", ["Model endpoint; internal collectors", "Optional cache; 8 TiB backup storage"]),
            ],
        },
        {
            "bounds": (665, 170, 1265, 1125), "title": "DR — WARM, SCALE TO FULL LOAD", "fill": "#EFF6F1", "outline": "#4A8062",
            "items": [
                ("Ingress / WAF", ["2 × 8 vCPU / 16 GiB", "Standby / active on failover"]),
                ("Gateway worker pool", ["4 × 32 vCPU / 128 GiB reserved", "3 pods warm → scale to 12"]),
                ("PostgreSQL replica site", ["3 × 32 vCPU / 256 GiB nodes", "4 TiB replicated dataset"]),
                ("Private dependencies", ["Private model; internal collectors", "Optional cache; 8 TiB backup storage"]),
            ],
        },
        {
            "bounds": (1280, 170, 1850, 1125), "title": "ISOLATED UAT — QUALIFICATION", "fill": "#F7F5EF", "outline": "#8A7446",
            "items": [
                ("Ingress / WAF", ["2 × 4 vCPU / 8 GiB", "Internal test access"]),
                ("Gateway worker pool", ["3 × 16 vCPU / 64 GiB nodes", "3 pods × 4 vCPU / 8 GiB"]),
                ("Separate PostgreSQL", ["3 × 16 vCPU / 64 GiB nodes", "1 TiB usable; no prod writes"]),
                ("Private test services", ["Synthetic / masked data", "Test model; 2 TiB backup storage"]),
            ],
        },
    ]

    boxes = []
    for site in sites:
        x0, y0, x1, y1 = site["bounds"]
        draw.rounded_rectangle(site["bounds"], radius=20, fill=site["fill"], outline=site["outline"], width=3)
        draw.text((x0 + 18, y0 + 16), site["title"], font=f_label, fill="#18304F")
        item_boxes = []
        y = y0 + 68
        for title, lines in site["items"]:
            box = (x0 + 18, y, x1 - 18, y + 174)
            draw.rounded_rectangle(box, radius=14, fill="white", outline=site["outline"], width=2)
            draw.text((box[0] + 14, box[1] + 12), title, font=f_label, fill="#18304F")
            for line_index, line in enumerate(lines):
                draw.text((box[0] + 14, box[1] + 58 + line_index * 42), line, font=f_body, fill="#293744")
            item_boxes.append(box)
            y += 198
        boxes.append(item_boxes)

    def arrow(start, end, color="#2A638C", dashed=False):
        x1, y1 = start
        x2, y2 = end
        if dashed:
            import math
            dx, dy = x2 - x1, y2 - y1
            length = max(1, (dx * dx + dy * dy) ** 0.5)
            ux, uy = dx / length, dy / length
            position = 0
            while position < length - 18:
                stop = min(position + 10, length - 18)
                draw.line((x1 + ux * position, y1 + uy * position, x1 + ux * stop, y1 + uy * stop), fill=color, width=4)
                position += 17
        else:
            draw.line((x1, y1, x2, y2), fill=color, width=4)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 15
        left = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
        right = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), left, right], fill=color)

    for item_boxes in boxes:
        for upper, lower in zip(item_boxes, item_boxes[1:]):
            arrow(((upper[0] + upper[2]) // 2, upper[3]), ((lower[0] + lower[2]) // 2, lower[1]))

    dc_db, dr_db = boxes[0][2], boxes[1][2]
    arrow((dc_db[2], (dc_db[1] + dc_db[3]) // 2), (dr_db[0], (dr_db[1] + dr_db[3]) // 2), color="#4A8062")
    draw.text((1299, 1147), "UAT isolated from production data and secrets.", font=f_note, fill="#8A7446")
    draw.text((50, 1182), "Green arrow = encrypted database replication over the bank WAN. Pod CPU/memory are part of the worker pools shown.", font=f_note, fill="#536273")
    draw.text((50, 1210), "Final replicas, storage and private-model capacity require workload qualification.", font=f_note, fill="#536273")
    image.save(path, quality=95)


def draw_browser_guard_flow(path):
    width, height = 2000, 1370
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f_title = ImageFont.truetype(str(font_path(True)), 38)
    f_label = ImageFont.truetype(str(font_path(True)), 24)
    f_body = ImageFont.truetype(str(font_path()), 20)
    f_small = ImageFont.truetype(str(font_path()), 18)
    draw.text((55, 28), "Browser AI Guard prompt, policy and telemetry flow", font=f_title, fill="#182635")

    def box(xy, title, lines, fill="#EAF1F7", outline="#35688F"):
        draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=3)
        x0, y0, x1, _ = xy
        draw.text((x0 + 16, y0 + 15), title, font=f_label, fill="#18304F")
        y = y0 + 55
        for line in lines:
            draw.text((x0 + 16, y), line, font=f_body, fill="#293744")
            y += 32

    def arrow(start, end, color="#2A638C", dashed=False):
        x1, y1 = start
        x2, y2 = end
        if dashed:
            dx, dy = x2 - x1, y2 - y1
            length = max(1, (dx * dx + dy * dy) ** 0.5)
            ux, uy = dx / length, dy / length
            pos = 0
            while pos < length - 20:
                endpos = min(pos + 12, length - 20)
                draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * endpos, y1 + uy * endpos), fill=color, width=4)
                pos += 19
        else:
            draw.line((x1, y1, x2, y2), fill=color, width=4)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 17
        left = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
        right = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), left, right], fill=color)

    def diamond(cx, cy, rx, ry, label, fill="#FFF4E5"):
        points = [(cx, cy - ry), (cx + rx, cy), (cx, cy + ry), (cx - rx, cy)]
        draw.polygon(points, fill=fill, outline="#8A7446")
        draw.line(points + [points[0]], fill="#8A7446", width=3)
        draw.multiline_text((cx, cy), label, font=f_label, fill="#18304F", anchor="mm", align="center", spacing=4)

    box((55, 150, 355, 330), "Employee / browser", ["Prompt or upload", "to an AI destination"], fill="#F2F6FA")
    box((435, 150, 800, 330), "YP-Raksha Guard", ["Extension / endpoint agent", "PAC and local proxy path"], fill="#F5F1F7", outline="#775A87")
    diamond(1000, 240, 155, 100, "Target site\nmonitored?")
    box((1260, 150, 1900, 330), "No target match", ["No Guard inspection for this site;", "continue under bank browser/network policy"], fill="#F7F5EF", outline="#8A7446")
    box((805, 420, 1195, 590), "Rule evaluation", ["Active configured rules", "Regex / content / optional AI rule"], fill="#EAF1F7")
    box((1300, 420, 1900, 590), "Optional AI Bot evaluation", ["Configured private model service", "Bot Answered when enabled"], fill="#F7F5EF", outline="#8A7446")
    diamond(1000, 700, 165, 95, "Policy\nviolation?")
    box((285, 820, 690, 1010), "BLOCK / SITE BLOCK", ["Show in-browser warning", "Stop submission to target"], fill="#FFF0EE", outline="#A13A36")
    box((790, 820, 1210, 1010), "ALLOW / REDACT / WARN", ["Apply configured action", "Forward only if policy permits"], fill="#EFF6F1", outline="#4A8062")
    box((1220, 820, 1510, 1010), "BOT ANSWERED", ["Return answer in browser", "Do not forward prompt"], fill="#EAF1F7", outline="#35688F")
    box((1540, 820, 1900, 1010), "AI destination", ["Approved/private endpoint", "Allowed request + response path"], fill="#EFF6F1", outline="#4A8062")

    draw.rounded_rectangle((630, 1080, 1935, 1300), radius=18, fill="#F6F9FC", outline="#18304F", width=4)
    draw.text((1185, 1090), "BANK-CONTROLLED INTERNAL EVENT AND ADMIN PLANE", font=f_label, fill="#18304F")
    box((675, 1145, 1000, 1270), "Gateway event API", ["POST /api/browser-ai/intercept", "internal HTTPS"], fill="#EAF1F7")
    box((1120, 1145, 1465, 1270), "LogsStore", ["browser_ai_logs", "browser_ai_search_logs*"], fill="#F7F5EF", outline="#8A7446")
    box((1580, 1145, 1895, 1270), "Browser AI UI", ["Authorized logs / insights", "admin review"], fill="#EFF6F1", outline="#4A8062")

    draw.text((1180, 195), "NO", font=f_small, fill="#8A7446")
    draw.text((1015, 350), "YES", font=f_small, fill="#2A638C")
    draw.text((1205, 470), "IF AI BOT RULE", font=f_small, fill="#8A7446")
    draw.text((735, 742), "YES: BLOCK", font=f_small, fill="#A13A36")
    draw.text((1170, 742), "NO: ALLOW / WARN / REDACT", font=f_small, fill="#4A8062")
    arrow((355, 240), (435, 240))
    arrow((800, 240), (845, 240))
    arrow((1155, 240), (1260, 240), color="#8A7446")
    arrow((1000, 340), (1000, 420))
    arrow((1195, 505), (1300, 505), color="#8A7446")
    arrow((1600, 590), (1365, 820), color="#8A7446")
    arrow((1000, 590), (1000, 605))
    arrow((835, 700), (690, 865), color="#A13A36")
    arrow((1165, 700), (1720, 820), color="#4A8062")
    draw.line((1210, 915, 1210, 790), fill="#4A8062", width=4)
    draw.line((1210, 790, 1500, 790), fill="#4A8062", width=4)
    arrow((1500, 790), (1540, 830), color="#4A8062")
    arrow((485, 1010), (790, 1145), color="#775A87", dashed=True)
    arrow((1000, 1010), (900, 1145), color="#775A87", dashed=True)
    arrow((1365, 1010), (1000, 1145), color="#775A87", dashed=True)
    arrow((1000, 1207), (1120, 1207))
    arrow((1465, 1207), (1580, 1207))
    draw.text((62, 1320), "* Monitored policy outcomes create internal events. Search records are separate where applicable. No-target traffic is not inspected; bank network policy must deny prohibited sites.", font=f_small, fill="#536273")
    image.save(path, quality=95)


def draw_logical_data_model(path):
    width, height = 1900, 1250
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    f_title = ImageFont.truetype(str(font_path(True)), 38)
    f_section = ImageFont.truetype(str(font_path(True)), 26)
    f_label = ImageFont.truetype(str(font_path(True)), 23)
    f_body = ImageFont.truetype(str(font_path()), 20)
    f_small = ImageFont.truetype(str(font_path()), 17)
    draw.text((55, 28), "Logical data model and record relationships", font=f_title, fill="#182635")

    columns = [
        ((55, 105, 625, 1160), "CONFIGURATION AND WORKSPACE", "#F6F9FC", "#35688F"),
        ((665, 105, 1215, 1160), "RUNTIME COMPONENTS", "#FAF8F2", "#8A7446"),
        ((1255, 105, 1845, 1160), "EVENTS AND DURABLE DATA", "#F4F8F5", "#4A8062"),
    ]
    for bounds, heading, fill, outline in columns:
        draw.rounded_rectangle(bounds, radius=22, fill=fill, outline=outline, width=4)
        draw.text((bounds[0] + 20, bounds[1] + 18), heading, font=f_section, fill="#18304F")

    def box(bounds, title, lines, fill="#FFFFFF", outline="#7C8B99"):
        draw.rounded_rectangle(bounds, radius=15, fill=fill, outline=outline, width=3)
        x0, y0, x1, _ = bounds
        draw.text((x0 + 18, y0 + 15), title, font=f_label, fill="#18304F")
        y = y0 + 55
        for line in lines:
            draw.multiline_text((x0 + 18, y), line, font=f_body, fill="#293744", spacing=5)
            y += 36 * (line.count("\n") + 1)

    def arrow(start, end, color="#2A638C", dashed=False):
        x1, y1 = start
        x2, y2 = end
        if dashed:
            import math
            dx, dy = x2 - x1, y2 - y1
            length = max(1, (dx * dx + dy * dy) ** 0.5)
            ux, uy = dx / length, dy / length
            pos = 0
            while pos < length - 22:
                endpos = min(pos + 12, length - 22)
                draw.line((x1 + ux * pos, y1 + uy * pos, x1 + ux * endpos, y1 + uy * endpos), fill=color, width=4)
                pos += 20
        else:
            draw.line((x1, y1, x2, y2), fill=color, width=4)
        import math
        angle = math.atan2(y2 - y1, x2 - x1)
        head = 16
        left = (x2 - head * math.cos(angle - 0.55), y2 - head * math.sin(angle - 0.55))
        right = (x2 - head * math.cos(angle + 0.55), y2 - head * math.sin(angle + 0.55))
        draw.polygon([(x2, y2), left, right], fill=color)

    box((82, 175, 598, 410), "Identity and access", [
        "governance_users  •  rbac_roles",
        "access_profiles  •  sessions",
        "section grants  •  audit_logs",
    ], fill="#FFFFFF", outline="#6C8295")
    box((82, 450, 598, 715), "Gateway configuration", [
        "providers / provider keys / models",
        "virtual keys / budgets / rate limits",
        "routing rules / plugins / feature flags",
    ], fill="#FFFFFF", outline="#6C8295")
    box((82, 755, 598, 1055), "Browser AI policy and fleet", [
        "targets / rules / interaction controls",
        "fleet defaults / agent inventory",
        "uninstall and package/update metadata",
    ], fill="#FFFFFF", outline="#6C8295")

    box((705, 270, 1175, 565), "Gateway control and inference", [
        "validates identity and request",
        "applies governance and policy",
        "routes to the selected provider",
        "writes configured operational records",
    ], fill="#EAF1F7", outline="#35688F")
    box((705, 745, 1175, 1010), "Guard agent and proxy", [
        "reads approved targets and rules",
        "evaluates supported browser traffic",
        "applies verdict and posts event",
    ], fill="#F5F1F7", outline="#775A87")

    box((1295, 175, 1805, 440), "Gateway and service records", [
        "inference / MCP / trace logs",
        "aggregates and configured metrics",
        "audit and operational history",
    ], fill="#FFFFFF", outline="#6C8295")
    box((1295, 510, 1805, 755), "Browser AI event records", [
        "browser_ai_logs",
        "browser_ai_search_logs",
        "prompt, target and agent metadata",
    ], fill="#FFFFFF", outline="#6C8295")
    box((1295, 825, 1805, 1065), "Attachment and diagnostic storage", [
        "attachment bytes: application data path",
        "metadata and expiry: relational record",
        "metrics / logs / traces: internal sinks",
    ], fill="#FFFFFF", outline="#6C8295")

    arrow((598, 295), (705, 345), color="#35688F")
    arrow((598, 585), (705, 475), color="#35688F")
    arrow((598, 900), (705, 875), color="#775A87")
    arrow((1175, 350), (1295, 305), color="#35688F")
    arrow((1175, 875), (1295, 640), color="#775A87")
    arrow((1550, 755), (1550, 825), color="#8A7446", dashed=True)
    draw.text((55, 1185), "Logical ownership and associations only; the selected release's migrations define the physical schema. Configure backup, retention and access per data class.", font=f_small, fill="#536273")
    image.save(path, quality=95)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.34)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Heading 1", 15, 15, 7), ("Heading 2", 12, 10, 5), ("Heading 3", 10.5, 8, 4)):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("000000")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title_style = styles["Title"]
    title_style.font.name = "Calibri"
    title_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title_style.font.size = Pt(27)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor.from_string("000000")
    title_style.paragraph_format.space_after = Pt(10)
    title_ppr = title_style._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.1)
    code.font.color.rgb = RGBColor.from_string("334155")
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0


def add_page_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=8, color=MUTED)


def add_running_header_footer(doc, title):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(f"YP-RAKSHA  |  {title}")
    set_font(run, size=8, bold=True, color=MUTED)

    def populate_footer(footer):
        table = footer.add_table(rows=1, cols=2, width=Inches(7.0))
        table.autofit = False
        table.columns[0].width = Inches(4.6)
        table.columns[1].width = Inches(2.4)
        set_table_borders(table)
        for idx, cell in enumerate(table.rows[0].cells):
            cell.width = Inches(4.6 if idx == 0 else 2.4)
            set_cell_margins(cell, 0, 0, 0, 0)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = table._tbl.tblPr.find(qn("w:tblBorders")).find(qn(f"w:{edge}"))
            if b is not None:
                b.set(qn("w:val"), "nil")
        left = table.cell(0, 0).paragraphs[0]
        left.paragraph_format.space_after = Pt(0)
        set_font(left.add_run("RAKSHA  •  YP-Raksha  •  Version 1.5"), size=8, color=MUTED)
        right = table.cell(0, 1).paragraphs[0]
        right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        right.paragraph_format.space_after = Pt(0)
        set_font(right.add_run("Page "), size=8, color=MUTED)
        add_page_field(right, "PAGE")
        set_font(right.add_run(" of "), size=8, color=MUTED)
        add_page_field(right, "NUMPAGES")

    populate_footer(section.footer)
    if section.different_first_page_header_footer:
        populate_footer(section.first_page_footer)


def parse_table_line(line):
    return [part.strip() for part in line.strip().strip("|").split("|")]


def add_markdown_table(doc, raw_rows):
    rows = [parse_table_line(line) for line in raw_rows]
    if len(rows) < 2:
        return
    header = rows[0]
    body = rows[2:]  # row 2 is the Markdown separator
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    sample = [header] + body
    weights = []
    for col in range(ncols):
        longest = max((len(row[col]) if col < len(row) else 0 for row in sample), default=12)
        weights.append(min(42, max(9, longest)))
    total = sum(weights) or ncols
    widths = [7.0 * weight / total for weight in weights]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)

    def populate(row_obj, values, header_row=False, row_index=0):
        for col, cell in enumerate(row_obj.cells):
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 90, 105, 90, 105)
            if header_row:
                set_cell_shading(cell, NAVY)
            elif row_index % 2:
                set_cell_shading(cell, PALE2)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            value = values[col] if col < len(values) else ""
            add_inline(p, value, base_size=8.2, base_color=("FFFFFF" if header_row else TEXT))
            for run in p.runs:
                run.bold = bool(header_row) or run.bold
                if header_row:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                if run.font.size is None:
                    run.font.size = Pt(8.2)

    populate(table.rows[0], header, header_row=True)
    set_cant_split(table.rows[0])
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True
    set_repeat_table_header(table.rows[0])
    for index, values in enumerate(body):
        row = table.add_row()
        set_cant_split(row)
        populate(row, values, row_index=index)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def collect_toc(source):
    entries = []
    for line in source.splitlines():
        match = re.match(r"^(#{2,3})\s+(.+?)\s*$", line)
        if match:
            depth = len(match.group(1)) - 1
            title = clean_heading(match.group(2))
            if title:
                entries.append((title, depth))
    return entries


def add_cover_and_toc(doc, spec, entries, page_map=None):
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(35)
    cover.paragraph_format.space_after = Pt(18)
    if LOGO.exists():
        cover.add_run().add_picture(str(LOGO), width=Inches(2.15))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(7)
    set_font(p.add_run("RAKSHA - Real-time AI Knowledge Screening & Hazard Audit"), size=13, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run("Version 1.5"), size=10, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    set_font(p.add_run("YP-Raksha"), size=16, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    set_font(p.add_run(spec["kind"]), size=10, bold=True, color=BLUE)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(spec["title"]), size=27, bold=True, color="000000")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run(spec["subtitle"]), size=14, color=MUTED)

    doc.add_page_break()
    p = doc.add_paragraph(style="Title")
    ppr = p._p.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run("Contents and document index"), size=20, bold=True, color="000000")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run("Section page numbers include the cover. Headings are listed in reading order."), size=9, color=MUTED)

    toc = doc.add_table(rows=1, cols=2)
    toc.alignment = WD_TABLE_ALIGNMENT.CENTER
    toc.autofit = False
    toc.columns[0].width = Inches(6.5)
    toc.columns[1].width = Inches(0.5)
    set_table_borders(toc)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = toc._tbl.tblPr.find(qn("w:tblBorders")).find(qn(f"w:{edge}"))
        if border is not None:
            border.set(qn("w:val"), "nil")

    def style_toc_cell(cell, width, label, header=False, depth=1):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, 20, 90, 20, 90)
        if header:
            set_cell_shading(cell, NAVY)
        elif depth == 1:
            set_cell_shading(cell, PALE2)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if width == 0.5:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif not header and depth == 2:
            p.paragraph_format.left_indent = Inches(0.2)
        set_font(p.add_run(label), size=8.5 if depth == 2 else 9, bold=header or depth == 1, color=("FFFFFF" if header else (NAVY if depth == 1 else TEXT)))

    style_toc_cell(toc.rows[0].cells[0], 6.5, "SECTION", header=True)
    style_toc_cell(toc.rows[0].cells[1], 0.5, "PAGE", header=True)
    set_repeat_table_header(toc.rows[0])
    set_cant_split(toc.rows[0])
    for title, depth in entries:
        row = toc.add_row()
        set_cant_split(row)
        style_toc_cell(row.cells[0], 6.5, title, depth=depth)
        style_toc_cell(row.cells[1], 0.5, str((page_map or {}).get(title, "—")), depth=depth)

    doc.add_page_break()


def add_picture_caption(doc, picture_path, caption, width=6.95):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(picture_path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(9)
    set_font(cp.add_run(caption), size=8.5, italic=True, color=MUTED)


def render_markdown_body(doc, source, spec, diagram_dir):
    lines = source.splitlines()
    title_skipped = False
    i = 0
    diagram_count = 0
    page_break_pending = False
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip() == "<!-- page-break -->":
            page_break_pending = True
            i += 1
            continue
        if line.startswith("# ") and not title_skipped:
            title_skipped = True
            i += 1
            continue
        if line.startswith("## "):
            heading_text = clean_heading(line[3:])
            p = doc.add_paragraph(style="Heading 1")
            if page_break_pending:
                p.paragraph_format.page_break_before = True
                page_break_pending = False
            add_inline(p, heading_text, base_size=15, base_color="000000")
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_paragraph(style="Heading 2")
            if page_break_pending:
                p.paragraph_format.page_break_before = True
                page_break_pending = False
            add_inline(p, clean_heading(line[4:]), base_size=12, base_color="000000")
            i += 1
            continue
        if line.startswith("#### "):
            p = doc.add_paragraph(style="Heading 3")
            if page_break_pending:
                p.paragraph_format.page_break_before = True
                page_break_pending = False
            add_inline(p, clean_heading(line[5:]), base_size=10.5, base_color="000000")
            i += 1
            continue
        if line.startswith("```"):
            language = line[3:].strip().casefold()
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            if language == "site-topology":
                diagram_count += 1
                add_picture_caption(doc, diagram_dir / "site-deployment-topology.png", f"Figure {diagram_count}. Reference DC, DR and UAT capacity pools inside the bank-controlled India boundary. Values are planning estimates, not measured product limits.")
                continue
            if language == "platform-architecture":
                diagram_count += 1
                add_picture_caption(doc, diagram_dir / "platform-architecture.png", f"Figure {diagram_count}. Full YP-Raksha client, control plane, data plane, integrations, persistence and internal observability architecture. Optional components depend on configuration; bank network controls enforce the residency boundary.")
                continue
            if language == "product-architecture":
                diagram_count += 1
                add_picture_caption(doc, diagram_dir / "logical-architecture.png", f"Figure {diagram_count}. Logical YP-Raksha architecture. Optional services and destinations depend on configuration; the bank must enforce the residency boundary.")
                continue
            if language == "flow-browser-ai":
                diagram_count += 1
                add_picture_caption(doc, diagram_dir / "browser-guard-flow.png", f"Figure {diagram_count}. Browser AI Guard request, verdict and internal event flow. Exact behavior depends on configured rules and the validated release.")
                continue
            if language == "logical-data-model":
                diagram_count += 1
                add_picture_caption(doc, diagram_dir / "logical-data-model.png", f"Figure {diagram_count}. Logical data groups and associations. This view is conceptual and does not assert physical foreign-key constraints.")
                continue
            if language == "flow-error-handling":
                diagram_count += 1
                flow = FLOW_DIAGRAMS[language]
                add_picture_caption(doc, diagram_dir / flow["file"], f"Figure {diagram_count}. {flow['title']}. Branches are selected by the classified failure; enabled behavior depends on release configuration.")
                continue
            if language in FLOW_DIAGRAMS:
                diagram_count += 1
                flow = FLOW_DIAGRAMS[language]
                add_picture_caption(doc, diagram_dir / flow["file"], f"Figure {diagram_count}. {flow['title']}. This is a logical workflow; enabled stages depend on deployment configuration.")
                continue
            if language == "mermaid":
                diagram_count += 1
                add_picture_caption(doc, diagram_dir / "sovereign-data-flow.png", f"Figure {diagram_count}. Intended DC, DR and UAT flows remain within the bank-controlled India boundary; red dashed arrows are denied at the network edge.")
                continue
            for code_line in block:
                p = doc.add_paragraph(style="Code Block")
                p.paragraph_format.left_indent = Inches(0.08)
                p.paragraph_format.right_indent = Inches(0.02)
                set_font(p.add_run(code_line or " "), name="Consolas", size=8.1, color="334155")
            continue
        if line.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            raw_rows = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                raw_rows.append(lines[i])
                i += 1
            if page_break_pending:
                p = doc.add_paragraph()
                p.paragraph_format.page_break_before = True
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(0)
                page_break_pending = False
            add_markdown_table(doc, raw_rows)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_together = True
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line), base_size=10.2)
            i += 1
            continue
        if re.match(r"^\s*\d+[.)]\s+", line):
            match = re.match(r"^\s*(\d+)[.)]\s+(.*)", line)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_together = True
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            add_inline(p, f"{match.group(1)}. {match.group(2)}", base_size=10.2)
            i += 1
            continue
        if re.match(r"^\*\*[^*]+:\*\*", line):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, line.strip(), base_size=9.7, base_color=MUTED)
            i += 1
            continue
        if line.strip() in ("---", "***", "___"):
            i += 1
            continue

        paragraph_lines = [re.sub(r"^\s*>\s?", "", line.strip())]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "```")) and not lines[i].lstrip().startswith("|") and not re.match(r"^\s*[-*]\s+", lines[i]) and not re.match(r"^\s*\d+[.)]\s+", lines[i]) and not re.match(r"^\*\*[^*]+:\*\*", lines[i]):
            paragraph_lines.append(re.sub(r"^\s*>\s?", "", lines[i].strip()))
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.keep_together = True
        add_inline(p, " ".join(paragraph_lines), base_size=10.4)


def build_docx(spec, page_map, output_path, diagram_dir):
    source = (ROOT / spec["source"]).read_text(encoding="utf-8-sig")
    doc = Document()
    configure_document(doc)
    add_running_header_footer(doc, spec["kind"].title())
    entries = collect_toc(source)
    add_cover_and_toc(doc, spec, entries, page_map)
    render_markdown_body(doc, source, spec, diagram_dir)
    doc.core_properties.title = spec["title"]
    doc.core_properties.subject = spec["subtitle"]
    doc.core_properties.author = "YP-Raksha Product Documentation"
    doc.core_properties.keywords = "YP-Raksha, Browser AI Guard, technical documentation, bank deployment"
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    doc.save(output_path)
    return entries


def render(docx_path, output_dir, word_script, pdftoppm):
    output_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = output_dir / f"{docx_path.stem}.pdf"
    result = subprocess.run(
        [
            str(Path("C:/Windows/System32/WindowsPowerShell/v1.0/powershell.exe")),
            "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(word_script),
            "-InputDocx", str(docx_path), "-OutputPdf", str(pdf_path),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode:
        raise RuntimeError(f"Word render failed for {docx_path.name}:\n{result.stdout}\n{result.stderr}")
    prefix = output_dir / docx_path.stem
    raster = subprocess.run(
        [str(pdftoppm), "-png", "-r", "120", str(pdf_path), str(prefix)],
        capture_output=True,
        text=True,
        check=False,
    )
    if raster.returncode:
        raise RuntimeError(f"PDF rasterization failed for {docx_path.name}:\n{raster.stdout}\n{raster.stderr}")
    return pdf_path


def page_map_from_pdf(pdf_path, entries):
    pages = [normalize(page.extract_text() or "") for page in PdfReader(str(pdf_path)).pages]
    if not entries:
        return {}
    first = normalize(entries[0][0])
    body_candidates = [i for i, text in enumerate(pages) if first and first in text]
    body_start = max(body_candidates, default=0)
    mapped = {}
    previous = body_start
    for title, _depth in entries:
        needle = normalize(title)
        hits = [i for i, text in enumerate(pages) if i >= previous and needle and needle in text]
        if not hits:
            hits = [i for i, text in enumerate(pages) if i >= body_start and needle and needle in text]
        if hits:
            page_index = hits[0]
            mapped[title] = page_index + 1
            previous = page_index
        else:
            mapped[title] = "—"
    return mapped


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--word-script", type=Path, default=ROOT / "render_word.ps1")
    parser.add_argument("--qa-dir", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, default=ROOT)
    parser.add_argument("--only", choices=[spec["stem"] for spec in DOCS], help="Build only the selected document.")
    args = parser.parse_args()
    output_dir = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    work = args.qa_dir / "build"
    diagrams = args.qa_dir / "diagrams"
    work.mkdir(parents=True, exist_ok=True)
    diagrams.mkdir(parents=True, exist_ok=True)
    pdftoppm = Path(sys.executable).parents[1] / "native" / "poppler" / "Library" / "bin" / "pdftoppm.exe"
    draw_architecture(diagrams / "logical-architecture.png", product=True)
    draw_architecture(diagrams / "sovereign-data-flow.png", product=False)
    draw_platform_architecture(diagrams / "platform-architecture.png")
    draw_site_topology(diagrams / "site-deployment-topology.png")
    draw_browser_guard_flow(diagrams / "browser-guard-flow.png")
    draw_error_handling_diagram(diagrams / "error-recovery-flow.png")
    draw_logical_data_model(diagrams / "logical-data-model.png")
    for flow in FLOW_DIAGRAMS.values():
        if flow["file"] == "error-recovery-flow.png":
            continue
        draw_flow_diagram(diagrams / flow["file"], flow["title"], flow["steps"], flow["note"])

    selected_docs = [spec for spec in DOCS if args.only is None or spec["stem"] == args.only]
    for spec in selected_docs:
        source = (ROOT / spec["source"]).read_text(encoding="utf-8-sig")
        entries = collect_toc(source)
        pre = work / f"{spec['stem']}-pre.docx"
        build_docx(spec, {}, pre, diagrams)
        pre_pdf = render(pre, work / f"{spec['stem']}-pre-render", args.word_script, pdftoppm)
        page_map = page_map_from_pdf(pre_pdf, entries)

        final_docx = output_dir / f"{spec['stem']}.docx"
        build_docx(spec, page_map, final_docx, diagrams)
        final_render = args.qa_dir / spec["stem"]
        final_pdf = render(final_docx, final_render, args.word_script, pdftoppm)
        shutil.copy2(final_pdf, output_dir / f"{spec['stem']}.pdf")
        print(f"{spec['stem']}: {len(PdfReader(str(final_pdf)).pages)} pages; TOC mapped {sum(v != '—' for v in page_map.values())}/{len(entries)} headings")


if __name__ == "__main__":
    main()

