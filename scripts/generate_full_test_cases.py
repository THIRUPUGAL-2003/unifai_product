#!/usr/bin/env python3
"""All-section test cases from the scope specification."""

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "docs" / "client-pack"))
from build_scope_pdf import REQS  # noqa: E402

OUT = ROOT / "UnifAI_Complete_Test_Cases.xlsx"

NAVY = "0F2744"
TEAL = "0E7490"
SLATE = "334155"
BAND = "E2E8F0"
BLUE = "F0F9FF"
WHITE = "FFFFFF"
AMBER = "FEF3C7"

thin = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
wrap = Alignment(wrap_text=True, vertical="top")
header_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
section_font = Font(name="Calibri", size=12, bold=True, color=WHITE)
title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
label_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
body_font = Font(name="Calibri", size=10, color="1E293B")

HEADERS = [
    "Test Case ID",
    "Section",
    "Requirement ID",
    "Test Scenario / Objective",
    "Actor",
    "Preconditions",
    "Step-by-Step Test Procedure",
    "Test Data / Inputs",
    "Expected Result (Pass Criteria)",
    "Actual Result (Tester Notes)",
    "Status",
    "Priority",
    "Defect / Bug ID",
    "Tester Name",
    "Execution Date",
]

# Extra cases that the requirement row does not already spell out as its own negative.
EXTRAS = [
    ("1. Access and accounts", "AUTH-01", [
        ["TC-AUTH-01-N1", "Sign in with a wrong password", "user", "Approved account anitha.user exists.",
         "1. Open /login.\n2. Enter the username.\n3. Enter a wrong password.\n4. Click Sign In.",
         "Username: anitha.user\nPassword: Wrong#Pass1",
         "Sign-in fails. The workspace does not open. The password is not revealed.", "P1 - High"],
        ["TC-AUTH-02-N1", "Sign up with passwords that do not match", "New employee", "Sign-up page is open. SMTP may already be tested.",
         "1. Open /signup.\n2. Fill username and email.\n3. Enter two different passwords.\n4. Submit.",
         "Username: anitha.user\nEmail: anitha.user@example.com\nPassword: Anitha#User24\nConfirm: Anitha#User25",
         "The form shows that the passwords do not match. No pending user is created.", "P1 - High"],
        ["TC-AUTH-02-N2", "Sign up with every field blank", "New employee", "Sign-up page is open.",
         "1. Leave username, email, password, and confirm password empty.\n2. Submit.",
         "All fields blank",
         "The form does not submit. Each required field asks to be filled.", "P1 - High"],
    ]),
    ("2. Governance", "GOV-06", [
        ["TC-GOV-06-N1", "API call with a missing virtual key", "Application", "SET-02 passed. Auth is required.",
         "1. POST /v1/chat/completions with no Authorization header.\n2. Read the response.",
         "Model: gpt-4.1-mini\nMessage: Hello\nAuthorization: absent",
         "The call is refused. No vendor request is made.", "P1 - High"],
        ["TC-GOV-02-N1", "Create a user with an empty username", "admin", "Signed in as priya.admin.",
         "1. Open Governance, Users, Create user.\n2. Leave username empty.\n3. Save.",
         "Username: blank\nEmail: anitha.user@example.com",
         "The user is not created. The username field shows it is required.", "P1 - High"],
    ]),
    ("3. Models, routing, and guardrails", "MOD-01", [
        ["TC-MOD-01-N1", "Save a provider with an empty vendor key", "admin", "Signed in as priya.admin. Provider needs a key.",
         "1. Open Models, Model Providers.\n2. Add OpenAI.\n3. Leave the API key empty.\n4. Save.",
         "Provider: OpenAI\nKey name: openai-prod-1\nAPI key: blank",
         "The provider is not saved as a working key. Model Catalog does not gain models from it.", "P1 - High"],
        ["TC-MOD-10-N1", "A browser prompt does not create a Guardrails LLM row as a Prompt Log", "admin", "MOD-10 rule is active. A laptop agent is also active.",
         "1. Send a card number through a virtual-key API call.\n2. Send the same text on a monitored website.\n3. Compare LLM Logs and Prompt Logs.",
         "API text: a test card number\nBrowser text: the same value",
         "The API hit is in LLM Logs. The browser hit is in Prompt Logs. The two logs are not the same row.", "P1 - High"],
    ]),
    ("4. Observability and settings", "SET-01", [
        ["TC-SET-01-N1", "Password reset before SMTP is saved", "admin then user", "SMTP is disabled.",
         "1. Leave SMTP disabled.\n2. As anitha.user, open Forgot password.\n3. Submit the email.",
         "Email: anitha.user@example.com\nSMTP enabled: No",
         "No reset code arrives. AUTH-04 is not failed as a product defect until SET-01 passes.", "P1 - High"],
    ]),
    ("5. Browser AI and the laptop", "BAI-02", [
        ["TC-BAI-02-N1", "A PAN on a site that is not a target is not blocked by the PAN rule", "Employee", "India PAN rule is active. example.com is not a target website.",
         "1. Open a site that is not in Target Websites.\n2. Type My PAN is ABCDE1234F and send.",
         "Domain: not in Target Websites\nText: My PAN is ABCDE1234F",
         "The browser rule does not fire. Prompt Logs does not show Blocked for that domain.", "P2 - Medium"],
        ["TC-BAI-15-N1", "Import a spreadsheet row whose rule type is ai_bot", "admin", "Signed in to Guard Rules.",
         "1. Prepare a CSV with Rule Type ai_bot.\n2. Import it.",
         "Rule Name: bot import\nRule Type: ai_bot\nAction: WARN",
         "The import is rejected. The rule is not created. A regex row in the same file format can still be imported.", "P1 - High"],
    ]),
    ("6. MCP, prompts, and skills", "MCP-03", [
        ["TC-MCP-03-N1", "Call a tool that was not ticked in the tool group", "Application", "MCP-01, MCP-02, and MCP-03 passed. mcp.tool_execution is on.",
         "1. Call a tool that is not search_docs with the Support virtual key.",
         "Key: Support bot — production\nTool: a tool left unticked",
         "The call is refused. MCP Logs does not show a success for that tool.", "P1 - High"],
        ["TC-PR-02-N1", "Request a prompt environment that was not deployed", "Application", "Only production version 1 is deployed.",
         "1. Call the API with header x-uf-prompt-environment: staging and the virtual key.",
         "Header: x-uf-prompt-environment: staging\nDeployed: production only",
         "The staging version is not returned. Production is unchanged.", "P2 - Medium"],
    ]),
]


def steps_for(pre, inputs, action):
    data_lines = [line.strip() for line in str(inputs).split("\n") if line.strip()]
    data = "\n".join(f"   - {line}" for line in data_lines) if data_lines else f"   - {inputs}"
    return (
        f"1. Confirm the precondition: {pre}\n"
        f"2. Enter the test data exactly, without changing names or values:\n{data}\n"
        f"3. Carry out this check: {action}\n"
        f"4. Compare what you see with every sentence in Expected Result."
    )


def priority_for(req_id):
    if req_id.startswith(("AUTH", "GOV-06", "GOV-07", "BAI-0", "BAI-1", "SET-01", "MOD-01", "MOD-10")):
        return "P1 - High"
    if req_id.startswith("NFR"):
        return "P1 - High"
    return "P2 - Medium"


def build_cases():
    extras_by_section = {}
    for section, _after, extras in EXTRAS:
        extras_by_section.setdefault(section, []).extend(extras)
    cases = []
    for title, rows in REQS:
        for row in rows:
            req_id, text, actor, pre, inputs, expected = row
            cases.append({
                "section": title,
                "values": [
                    f"TC-{req_id}",
                    title.split(". ", 1)[-1],
                    req_id,
                    text,
                    actor,
                    pre,
                    steps_for(pre, inputs, text),
                    inputs,
                    expected,
                    "",
                    "Untested",
                    priority_for(req_id),
                    "",
                    "",
                    "",
                ],
            })
        for extra in extras_by_section.get(title, []):
            tc_id, text, actor, pre, procedure, data, expected, priority = extra
            cases.append({
                "section": title,
                "values": [
                    tc_id,
                    title.split(". ", 1)[-1],
                    tc_id.split("-N")[0].replace("TC-", ""),
                    text,
                    actor,
                    pre,
                    procedure,
                    data,
                    expected,
                    "",
                    "Untested",
                    priority,
                    "",
                    "",
                    "",
                ],
            })
    nfr = [
        ["TC-NFR-01", "Non-functional", "NFR-01", "Session cookie is HTTP-only after sign-in.", "admin", "AUTH-01 can pass.",
         "1. Sign in as priya.admin.\n2. Inspect the session cookie named token.",
         "Account: priya.admin", "The cookie is HTTP-only. On HTTPS it is also Secure.", "", "Untested", "P1 - High", "", "", ""],
        ["TC-NFR-02", "Non-functional", "NFR-02", "A user screen does not show a vendor API key.", "user", "A provider key exists. User anitha.user has no Settings access.",
         "1. Sign in as anitha.user.\n2. Open every granted page.\n3. Search the page text for the vendor key.",
         "User: anitha.user", "The vendor key is not visible.", "", "Untested", "P1 - High", "", "", ""],
        ["TC-NFR-03", "Non-functional", "NFR-03", "A user cannot open another team's prompt folder.", "user", "Two folders exist. This user is limited to Support.",
         "1. Sign in as anitha.user.\n2. Open Prompt Repository.",
         "Allowed folder: Support\nHidden folder: Finance", "Only Support is listed.", "", "Untested", "P2 - Medium", "", "", ""],
        ["TC-NFR-04", "Non-functional", "NFR-04", "Guard data is per Windows user.", "IT", "Gateway Guard installed.",
         "1. Sign in to Windows as user A and note %LOCALAPPDATA%\\Gateway\\Guard.\n2. Sign in as user B and open that same path.",
         "Path: %LOCALAPPDATA%\\Gateway\\Guard", "Each Windows user has a separate folder.", "", "Untested", "P2 - Medium", "", "", ""],
        ["TC-NFR-05", "Non-functional", "NFR-05", "One alert incident is sent at most once in 10 minutes.", "admin", "OBS-05 test message already succeeded.",
         "1. Hit the same budget twice within 10 minutes.\n2. Count alert messages.",
         "Channel: Ops Slack", "One notice arrives for that incident inside 10 minutes.", "", "Untested", "P1 - High", "", "", ""],
        ["TC-NFR-06", "Non-functional", "NFR-06", "Server health endpoint succeeds when the service is up.", "IT", "UnifAI is running.",
         "1. Call GET /health.",
         "URL: https://unifai.example.com/health", "The response is success.", "", "Untested", "P1 - High", "", "", ""],
    ]
    for row in nfr:
        cases.append({"section": "7. Non-functional", "values": row})
    return cases


def style_header(ws):
    for col, name in enumerate(HEADERS, 1):
        cell = ws.cell(1, col, name)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O1"
    widths = [18, 28, 16, 42, 22, 36, 55, 40, 55, 28, 14, 14, 16, 16, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_cases(ws, cases):
    style_header(ws)
    row = 2
    current = None
    fills = [PatternFill("solid", fgColor=BLUE), PatternFill("solid", fgColor=WHITE)]
    data_i = 0
    for item in cases:
        if item["section"] != current:
            current = item["section"]
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=len(HEADERS))
            cell = ws.cell(row, 1, current.upper())
            cell.font = section_font
            cell.fill = PatternFill("solid", fgColor=SLATE)
            cell.alignment = Alignment(vertical="center")
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row, col).fill = PatternFill("solid", fgColor=SLATE)
                ws.cell(row, col).border = thin
            ws.row_dimensions[row].height = 22
            row += 1
        for col, value in enumerate(item["values"], 1):
            cell = ws.cell(row, col, value)
            cell.font = body_font
            cell.alignment = wrap
            cell.border = thin
            cell.fill = fills[data_i % 2]
        texts = [str(v) for v in item["values"]]
        widths = [18, 28, 16, 42, 22, 36, 55, 40, 55, 28, 14, 14, 16, 16, 16]

        def wrapped_lines(text, width):
            total = 0
            for part in text.split("\n"):
                total += max(1, (len(part) + width - 1) // width)
            return total

        lines = max(wrapped_lines(text, width) for text, width in zip(texts, widths))
        ws.row_dimensions[row].height = lines * 16 + 10
        data_i += 1
        row += 1
    last = row - 1
    dv = DataValidation(type="list", formula1='"Untested,Pass,Fail,Blocked"', allow_blank=False)
    dv.error = "Choose Untested, Pass, Fail, or Blocked"
    dv.errorTitle = "Status"
    dv.add(f"K2:K{last}")
    ws.add_data_validation(dv)
    ws.conditional_formatting.add(f"K2:K{last}", FormulaRule(formula=['$K2="Pass"'], fill=PatternFill("solid", fgColor="BBF7D0")))
    ws.conditional_formatting.add(f"K2:K{last}", FormulaRule(formula=['$K2="Fail"'], fill=PatternFill("solid", fgColor="FECACA")))
    ws.conditional_formatting.add(f"K2:K{last}", FormulaRule(formula=['$K2="Blocked"'], fill=PatternFill("solid", fgColor=AMBER)))
    ws.auto_filter.ref = f"A1:O{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddHeader.left.text = "UnifAI complete test cases"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.print_title_rows = "1:1"
    ws.page_setup.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    return last


def main():
    cases = build_cases()
    wb = openpyxl.Workbook()
    plan = wb.active
    plan.title = "QA Test Plan"
    plan["B2"] = "UnifAI — complete test cases for every section"
    plan["B2"].font = title_font
    plan["B3"] = "Taken from the scope specification. One requirement ID, one test case, plus the negative cases marked -N."
    plan["B3"].font = Font(name="Calibri", size=11, italic=True, color="475569")
    plan.column_dimensions["B"].width = 36
    plan.column_dimensions["C"].width = 88
    info = [
        ("Product", "UnifAI website, Gateway Guard laptop agent, YP-Raksha extension"),
        ("Specification", "docs/client-pack/07_Scope_Specification.pdf and docs/SCOPE_SPECIFICATION.md"),
        ("How to execute", "Run IDs in the order in the scope. Do not fail a later test while the earlier ID is still failing."),
        ("Sample admin", "priya.admin / Priya#Admin24"),
        ("Sample sub_admin", "ravi.security / Ravi#Secure24"),
        ("Sample user", "anitha.user / Anitha#User24"),
        ("Sample website", "https://unifai.example.com"),
        ("Status values", "Untested, Pass, Fail, Blocked"),
        ("Total test cases", str(len(cases))),
    ]
    r = 5
    for label, value in info:
        plan.cell(r, 2, label).font = label_font
        plan.cell(r, 2).fill = PatternFill("solid", fgColor=BAND)
        plan.cell(r, 2).border = thin
        plan.cell(r, 3, value).font = body_font
        plan.cell(r, 3).border = thin
        r += 1

    counts = {}
    for item in cases:
        counts[item["section"]] = counts.get(item["section"], 0) + 1
    r += 1
    plan.cell(r, 2, "Section").font = header_font
    plan.cell(r, 2).fill = PatternFill("solid", fgColor=NAVY)
    plan.cell(r, 3, "Test cases").font = header_font
    plan.cell(r, 3).fill = PatternFill("solid", fgColor=NAVY)
    first_count = r + 1
    for section, count in counts.items():
        r += 1
        plan.cell(r, 2, section).font = body_font
        plan.cell(r, 2).border = thin
        plan.cell(r, 3, count).font = body_font
        plan.cell(r, 3).border = thin
    last_count = r
    chart = BarChart()
    chart.type = "col"
    chart.title = "Test cases by section"
    chart.y_axis.title = "Cases"
    chart.style = 10
    data = Reference(plan, min_col=3, min_row=first_count - 1, max_row=last_count)
    cats = Reference(plan, min_col=2, min_row=first_count, max_row=last_count)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.y_axis.majorGridlines = None
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True
    chart.width = 18
    chart.height = 8
    plan.add_chart(chart, "B16")

    tests = wb.create_sheet("All Section Test Cases")
    add_cases(tests, cases)

    bugs = wb.create_sheet("Defect Log")
    bug_headers = ["Defect ID", "Linked Test Case ID", "Section", "Summary", "Severity", "Steps", "Observed", "Expected", "Reported By", "Date", "Status"]
    for col, name in enumerate(bug_headers, 1):
        cell = bugs.cell(1, col, name)
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor=TEAL)
        bugs.column_dimensions[get_column_letter(col)].width = 24
    bugs.freeze_panes = "A2"
    bugs.auto_filter.ref = "A1:K1"

    wb.save(OUT)
    print(OUT.name, len(cases))


if __name__ == "__main__":
    main()
